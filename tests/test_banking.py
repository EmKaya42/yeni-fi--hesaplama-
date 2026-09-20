import io
import json
import time
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from test_workflow import RECEIPT, Z_REPORT, client, profile, server, successful_read, upload
from services.account_chart import read_chart, select_account
from services.accounting import journal_rows
from services.banking import BANK_BY_CODE, BANKS, bank_from_iban, identify_banks, normalize_iban
from services.document_extraction import extract_document
from services.document_ocr import RetryableOCRError
from services.queue_worker import process_next
from services.storage import database, initialize


def test_iban(code, suffix="0000000000000001"):
    bban = code.zfill(5) + "0" + suffix
    check = 98 - int(bban + "292700") % 97
    return f"TR{check:02}{bban}"
test_iban.__test__ = False


def chart_csv(extra=""):
    return ("Hesap Kodu;Hesap Adı;Banka Kodu;IBAN;Kart Son 4 Hane\n"
            "770;Genel Yönetim Giderleri;;;\n770.001;Kırtasiye Giderleri;;;\n"
            "191;İndirilecek KDV;;;\n191.20;%20 İndirilecek KDV;;;\n"
            "191.10;%10 İndirilecek KDV;;;\n100.01;Ana Kasa;;;\n"
            "600.20;%20 KDV'li Satış;;;\n391.20;%20 Hesaplanan KDV;;;\n"
            f"102.1;Garanti BBVA Banka Hesabı;0062;{test_iban('0062')};4321\n"
            f"102.10;İş Bankası Banka Hesabı;0064;{test_iban('0064')};8765\n"
            "309.01;Garanti Kredi Kartı;0062;;1234\n309.02;Akbank Kredi Kartı;0046;;5678\n"
            "108.01;Akbank POS Tahsilat;0046;;\n108.02;Garanti POS Tahsilat;0062;;\n" + extra)


def import_chart(client, name="TEST firma.csv", chart_id="", contents=None, headers=None):
    response = client.post("/api/account-charts", headers=headers or client.auth,
                           data={"program":"luca", "chart_id":chart_id, "file":(io.BytesIO((contents or chart_csv()).encode()), name)})
    assert response.status_code == 201, response.json
    return response.json["chart"]["id"]


def document(payment, kind="receipts", direction="expense"):
    text = RECEIPT.replace("NAKIT 120,00", payment) if kind == "receipts" else Z_REPORT
    data = extract_document(text, kind)
    return {"kind":kind, "direction":direction, "filename":"TEST.png", "result":data}


def configured_chart():
    return {**profile(), "chart":{"name":"TEST", "accounts":read_chart(io.BytesIO(chart_csv().encode()), "test.csv")}}


def test_bank_catalog_includes_current_2026_participants():
    assert len(BANKS) == len(BANK_BY_CODE) == 71
    assert BANK_BY_CODE["0215"]["name"] == "Adil Katılım"
    assert BANK_BY_CODE["0216"]["name"] == "İktisat Katılım"
    assert BANK_BY_CODE["0137"]["name"] == "Hepsi Bank"
    assert identify_banks("GARANTİ BBVA") == ["0062"]
    assert identify_banks("ZİRAAT KATILIM") == ["0209"]
    assert identify_banks("VAKIF KATILIM") == ["0210"]
    assert identify_banks("GARANTİLİ ÜRÜN / BANKA") == []


def test_iban_checksum_and_bank_code_are_separate_from_ledger():
    assert normalize_iban("TR47 0000 1001 0000 0350 9300 01") == "TR470000100100000350930001"
    assert bank_from_iban(test_iban("0062")) == "0062"
    with pytest.raises(ValueError):
        normalize_iban("TR480000100100000350930001")


@pytest.mark.parametrize("payment,account", [("NAKIT 120,00", "100"), ("KREDI KARTI 120,00", "309"), ("BANKA KARTI 120,00", "102"), ("HAVALE 120,00", "102"), ("EFT 120,00", "102"), ("FAST 120,00", "102")])
def test_payment_methods_use_correct_standard_accounts(payment, account):
    doc = document(payment)
    assert doc["result"]["issues"] == []
    rows = journal_rows([doc], profile())
    assert rows[-1]["account"] == account
    assert sum((row["debit"]-row["credit"] for row in rows), Decimal(0)) == 0


def test_partial_cash_and_card_payment_is_automatic_and_balanced():
    doc = document("NAKIT 50,00\nKREDI KARTI 70,00")
    assert not doc["result"]["issues"]
    assert doc["result"]["payment_method"] == "mixed"
    rows = journal_rows([doc], profile())
    assert [(row["account"], row["credit"]) for row in rows[-2:]] == [("100", Decimal("50.00")), ("309", Decimal("70.00"))]


def test_generic_pos_does_not_imply_business_credit_card():
    doc = document("POS 120,00")
    assert doc["result"]["issues"]
    with pytest.raises(ValueError):
        journal_rows([doc], profile())


def test_missing_payment_is_not_posted_to_arbitrary_supplier():
    assert document("")["result"]["issues"]


def test_pos_acquirer_is_not_card_issuer():
    doc = document("AKBANK KREDI KARTI 120,00\nKART NO ****1234")
    assert not doc["result"]["issues"]
    rows = journal_rows([doc], configured_chart())
    assert rows[-1]["account"] == "309.01"  # Garanti card, Akbank POS.
    assert rows[-1]["bank_code"] == "0062"
    assert rows[0]["account"] == "770.001" and rows[1]["account"] == "191.20"


def test_issuer_conflict_with_last_four_blocks_posting():
    doc = document("KREDI KARTI 120,00\nKART NO ****1234\nKARTI VEREN BANKA: AKBANK")
    with pytest.raises(ValueError, match="eşleşmedi"):
        journal_rows([doc], configured_chart())


def test_two_card_accounts_without_identity_are_not_guessed():
    with pytest.raises(ValueError, match="birden fazla"):
        journal_rows([document("KREDI KARTI 120,00")], configured_chart())


def test_debit_card_last_four_chooses_bank_account_not_credit_account():
    rows = journal_rows([document("BANKA KARTI 120,00\nKART NO ****8765")], configured_chart())
    assert rows[-1]["account"] == "102.10"


def test_sender_iban_selects_paying_account_and_recipient_is_not_used():
    doc = document(f"HAVALE 120,00\nGONDEREN IBAN: {test_iban('0062')}\nALICI IBAN: {test_iban('0064')}")
    assert not doc["result"]["issues"]
    assert journal_rows([doc], configured_chart())[-1]["account"] == "102.1"
    doc["direction"] = "income"
    assert journal_rows([doc], configured_chart())[-1]["account"] == "102.10"
    unknown = document(f"EFT 120,00\nALICI IBAN: {test_iban('0064')}")
    with pytest.raises(ValueError, match="birden fazla"):
        journal_rows([unknown], configured_chart())


def test_invalid_iban_requires_review():
    doc = document("HAVALE 120,00\nGONDEREN IBAN: TR480000100100000350930001")
    assert any("IBAN" in issue for issue in doc["result"]["issues"])


def test_z_bank_breakdown_does_not_double_count_card_summary():
    text = Z_REPORT.replace("KREDI KARTI 70,00", "KREDI KARTI 70,00\nAKBANK POS 30,00\nGARANTI POS 40,00")
    data = extract_document(text, "z-reports")
    assert not data["issues"]
    doc = {"kind":"z-reports", "direction":"income", "filename":"z.png", "result":data}
    rows = journal_rows([doc], configured_chart())
    assert [(row["account"], row["debit"]) for row in rows[-3:]] == [("100.01", Decimal(50)), ("108.01", Decimal(30)), ("108.02", Decimal(40))]
    assert extract_document(text.replace("GARANTI POS 40,00", "GARANTI POS 41,00"), "z-reports")["issues"]


def test_chart_parent_rows_and_sibling_codes_are_distinguished():
    accounts = configured_chart()["chart"]["accounts"]
    assert not next(account for account in accounts if account["code"] == "770")["leaf"]
    assert all(next(account for account in accounts if account["code"] == code)["leaf"] for code in ("102.1", "102.10"))


@pytest.mark.parametrize("extra", ["102.99;Akbank EUR;0046;;\n", "300.01;Garanti Banka Kredisi;0062;;\n"])
def test_foreign_currency_and_bank_loans_are_not_card_defaults(extra):
    accounts = read_chart(io.BytesIO(chart_csv(extra).encode()), "test.csv")
    row = accounts[-1]
    assert row["currency"] == "EUR" or row["role"] != "expense_card"


def test_chart_duplicate_codes_and_conflicting_iban_are_rejected():
    with pytest.raises(ValueError):
        read_chart(io.BytesIO(chart_csv("100.01;Başka Kasa;;;\n").encode()), "test.csv")
    with pytest.raises(ValueError):
        read_chart(io.BytesIO(chart_csv().replace(test_iban("0062"), test_iban("0064")).encode()), "test.csv")


def test_company_isolation_in_upload_deduplication_counts_and_export(client):
    first_chart = import_chart(client, "Firma A.csv")
    first = upload(client)
    process_next(server.DB_PATH, successful_read)
    second_chart = import_chart(client, "Firma B.csv")
    assert first_chart != second_chart
    assert client.get("/api/documents", headers=client.auth).json["items"] == []
    second = upload(client)  # Same bytes belong independently to this company.
    assert second != first
    process_next(server.DB_PATH, successful_read)
    assert client.get(f"/api/documents/{second}", headers=client.auth).json["status"] == "success"
    output = client.get("/export/excel", headers=client.auth)
    assert output.status_code == 200
    assert load_workbook(io.BytesIO(output.data)).active["A2"].value == "770.001"
    listing = client.get(f"/api/documents?program=luca&chart_id={first_chart}", headers=client.auth).json
    assert len(listing["items"]) == 1 and listing["items"][0]["id"] == first
    with database(server.DB_PATH) as db:
        assert db.execute("SELECT export_count FROM documents WHERE id=?", (first,)).fetchone()[0] == 0
        assert db.execute("SELECT export_count FROM documents WHERE id=?", (second,)).fetchone()[0] == 1


def test_imported_chart_ids_and_codes_cannot_be_forged_by_other_user(client):
    chart_id = import_chart(client)
    bob = client.headers_for("bob")
    assert client.get("/api/settings", headers=bob).json["charts"] == []
    assert client.put("/api/settings", headers=bob, json={"program":"luca", "chart_id":chart_id}).status_code == 400
    assert client.post("/api/account-charts", headers=bob, data={"program":"luca", "chart_id":chart_id, "file":(io.BytesIO(chart_csv().encode()), "stolen.csv")}).status_code == 400
    tampered = {"program":"luca", "chart_id":chart_id, "accounts":{"expense":"999"}, "chart":{"accounts":[{"code":"999"}]}}
    assert client.put("/api/settings", headers=client.auth, json=tampered).status_code == 200
    upload(client)
    process_next(server.DB_PATH, successful_read)
    output = client.get("/export/excel", headers=client.auth)
    assert load_workbook(io.BytesIO(output.data)).active["A2"].value == "770.001"


def test_chart_update_recalculates_accounting_without_repeating_ocr(client):
    chart_id = import_chart(client, contents=chart_csv("100.02;İkinci Kasa;;;\n"))
    doc_id = upload(client)
    process_next(server.DB_PATH, successful_read)
    assert client.get("/api/documents", headers=client.auth).json["counts"] == {"mapping":1}
    assert client.get("/export/excel", headers=client.auth).status_code == 400
    import_chart(client, chart_id=chart_id)
    doc = client.get(f"/api/documents/{doc_id}", headers=client.auth).json
    assert doc["status"] == "success" and doc["attempts"] == 1
    assert doc["accounting"]["rows"][-1]["account"] == "100.01"


def test_period_filter_search_and_export_history(client):
    one = upload(client)
    process_next(server.DB_PATH, successful_read)
    two = upload(client, 1)
    process_next(server.DB_PATH, lambda *args: extract_document(RECEIPT.replace("13.09.2026", "13.08.2026"), "receipts"))
    listing = client.get("/api/documents?period=2026-08&search=DEFTER", headers=client.auth).json
    assert listing["total"] == 1 and listing["items"][0]["id"] == two
    assert listing["periods"] == ["2026-09", "2026-08"]
    assert client.get("/export/excel?period=2026-08", headers=client.auth).status_code == 200
    assert client.get(f"/api/documents/{one}", headers=client.auth).json["export_count"] == 0
    assert client.get(f"/api/documents/{two}", headers=client.auth).json["export_count"] == 1
    with database(server.DB_PATH) as db:
        row = db.execute("SELECT * FROM export_log").fetchone()
        assert json.loads(row["document_ids"]) == [two] and len(row["digest"]) == 64


def test_period_filter_preserves_pending_queue_and_blocks_partial_export(client):
    upload(client)
    process_next(server.DB_PATH, successful_read)
    upload(client, 1)
    listing = client.get("/api/documents?period=2026-09", headers=client.auth).json
    assert listing["total"] == 1 and listing["counts"] == {"success": 1}
    assert listing["pending"] == 1 and listing["queue_counts"] == {"queued": 1}
    assert client.get("/export/excel?period=2026-09", headers=client.auth).status_code == 409


@pytest.mark.parametrize("name,currency", [("Garanti EUR Hesabı", "TRY"), ("Garanti EUR USD Hesabı", "EUR")])
def test_chart_rejects_conflicting_currency_evidence(name, currency):
    content = f"Hesap Kodu;Hesap Adı;Para Birimi\n102.01;{name};{currency}\n".encode("utf-8")
    with pytest.raises(ValueError, match="para birimi çelişiyor"):
        read_chart(io.BytesIO(content), "firma.csv")


def test_transient_ocr_retries_are_delayed_bounded_and_persisted(client, monkeypatch):
    doc_id = upload(client)
    current_time = time.time()
    monkeypatch.setattr("services.queue_worker.time.time", lambda: current_time)
    def timeout(*args):
        raise RetryableOCRError("timeout")
    for attempt in range(3):
        assert process_next(server.DB_PATH, timeout)
        initialize(server.DB_PATH)
        doc = client.get(f"/api/documents/{doc_id}", headers=client.auth).json
        assert doc["status"] == ("failed" if attempt == 2 else "queued")
        if attempt < 2:
            assert not process_next(server.DB_PATH, timeout)
        current_time += 20
    assert doc["attempts"] == 3 and doc["auto_retries"] == 2


def test_low_confidence_ocr_gets_one_automatic_alternative_read(client, monkeypatch):
    doc_id = upload(client)
    result = successful_read(None, 0, "receipts", 1)
    result.update(engine="Tesseract · tur+eng", issues=["İki okuma eşleşmedi."])
    process_next(server.DB_PATH, lambda *args: result)
    doc = client.get(f"/api/documents/{doc_id}", headers=client.auth).json
    assert doc["status"] == "queued" and doc["auto_retries"] == 1
    with database(server.DB_PATH) as db:
        db.execute("UPDATE documents SET retry_after=0 WHERE id=?", (doc_id,))
    process_next(server.DB_PATH, lambda *args: result)
    assert client.get(f"/api/documents/{doc_id}", headers=client.auth).json["status"] == "review"


def test_z_report_fallback_payments_from_belge_tipleri():
    z_text = (
        "ORNEK MARKET\n"
        "VERGİ DAİRESİ: Kadıköy\n"
        "VKN: 1234567890\n"
        "Z RAPORU NO: 000456\n"
        "MALİ SİCİL NO: AB00000123\n"
        "FİŞ ADEDİ: 2\n"
        "TARIH: 13.09.2026 SAAT: 23:15:42\n"
        "KDV %20 100,00 20,00\n"
        "SATIS TOPLAMI 120,00\n"
        "ODEME BILGILERI\n"
        "OKUNAMAYAN_ODEME\n"
        "BELGE TIPLERI\n"
        "-NAKIT *50,00\n"
        "-KREDI *70,00\n"
        "-DIGER *0,00\n"
        "SAYACLAR\n"
    )
    data = extract_document(z_text, "z-reports")
    assert not data["issues"], data["issues"]
    assert data["payment_entries"] == [
        {"method": "cash", "amount": "50.00", "bank_code": "", "bank_role": "unspecified"},
        {"method": "card", "amount": "70.00", "bank_code": "", "bank_role": "acquirer"},
        {"method": "other", "amount": "0.00", "bank_code": "", "bank_role": "unspecified"},
    ]
    assert data["cash_amount"] == "50.00"
    assert data["card_amount"] == "70.00"


def test_z_report_corrupted_totals_and_identifiers_require_review():
    raw_ocr = (
        "FORA TURIZM REKLAM\n"
        "SAR. TIC LTO STI\n"
        "SISLI V.D. 3880097945\n"
        "TARIH 30/05/2026\n"
        "SAAT 04:21:42\n"
        "Z RAPORU\n"
        "RAPOR NO 880\n"
        "MALI BELLEK TOPLAMI *12.867.454,44\n"
        "MALI BELLEK TOP KDV 2.070.030,56\n"
        "GMLK FS D#KHg\n"
        "TOPLAM *35.650,00\n"
        "Fopxdy *5.941,66\n"
        "~KDV BLGLER-\n"
        "Kdv  220.00.941,66\n"
        "TOPLAM 755.650,00\n"
        "~DEPARTHAN BLGLER\n"
        "BRA %20\n"
        "TOPLAM *34,400,00\n"
        "YERL K %20\n"
        "TupLak *1.250,00\n"
        "DEHE BLGLER\n"
        "NAKIT\n"
        "TOPLAM 0,00\n"
        "KREDI 33\n"
        "TOPLAM 435.650,00\n"
        "'~BELGE TIPLERI\n"
        "33\n"
        "KC FLERI\n"
        "TOPKDV *5.941,66\n"
        "-SATI ToplAMI *35.650,00\n"
        "-NAKIT *0,00\n"
        "KREDI 35.650,00\n"
        "-DGER 40,00\n"
        "5\n"
        "IPTAL\n"
        "~TOPKDV *775,00\n"
        "SATI   ToplAhi *4.650,00\n"
        "SAYACLAR -\n"
        "HAL F Adet 34\n"
        "33\n"
        "HTER F ADETI\n"
        "-KASYER BLG--\n"
        "*35.650,00\n"
        "KASIYERI\n"
        "KASIYER: KASIYER1\n"
        "EK NO:0001 Z NO: 1880\n"
        "NF JH 20004135\n"
    )
    data = extract_document(raw_ocr, "z-reports")
    assert data["issues"]
    assert any("çelişiyor" in issue or "birden fazla numara" in issue for issue in data["issues"])
    assert data["document_no"] == "880"  # Never substitute the different footer number.
    assert data["card_amount"] == "435650.00"  # Never trim digits to force agreement.


def test_z_report_zero_cash_does_not_invent_unreadable_card_amount():
    # Degraded OCR where KREDI amount was corrupted by checkmark or noise but Nakit is 0.00
    degraded = (
        "FORA TURIZM REKLAM\n"
        "SAR. TIC LTO STI\n"
        "SISLI V.D. 3880097945\n"
        "TARIH 30/05/2026\n"
        "SAAT 04:21:42\n"
        "Z RAPORU\n"
        "RAPOR NO 1880\n"
        "MALI BELLEK TOPLAMI *12.867.454,44\n"
        "MALI BELLEK TOP KDV 2.070.030,56\n"
        "GUNLUK FIS DOKUMU\n"
        "TOPLAM *35.650,00\n"
        "TOPKDV *5.941,66\n"
        "~KDV BILGILERI-\n"
        "KDV %20.00 *5.941,66\n"
        "TOPLAM *35.650,00\n"
        "~DEPARTMAN BILGILERI\n"
        "BIRA %20.00\n"
        "TOPLAM *34.400,00\n"
        "MIKTARI 32,0000\n"
        "YERLI ICKI %20.00\n"
        "TOPLAM *1.250,00\n"
        "MIKTARI 1,0000\n"
        "ODEME BILGILERI\n"
        "NAKIT\n"
        "TOPLAM 0,00\n"
        "KREDI 33\n"
        "Topla /55.650,C\n"
        "BELGE TIPLERI\n"
        "'Jii *0,00\n"
        "'Nil *35.650,00\n"
        "SAYACLAR\n"
        "MALI FIS ADET 34\n"
        "MUSTERI FISI ADETI 33\n"
        "KASIYER BILGI\n"
        "KASIYER1 *35.650,00\n"
        "KASIYER: KASIYER1\n"
        "EK NO:0001 Z NO: 1880\n"
        "JH 20004135\n"
    )
    data = extract_document(degraded, "z-reports")
    assert data["issues"]
    assert data["total_amount"] == "35650.00"
    assert data["card_amount"] == ""
    assert not any(e["method"] == "card" for e in data["payment_entries"])




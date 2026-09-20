import io
import json
import os
import tempfile
import time
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import rsa
from openpyxl import Workbook, load_workbook
from PIL import Image

os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="receipt-test-import-")
import app as server
from services.accounting import PROGRAMS, STANDARD_ACCOUNTS, default_profile, export_workbook, journal_rows, read_template, validate_profile
from services.document_extraction import decimal_money, extract_document
from services.document_ocr import read_tesseract_document as read_document
from services.queue_worker import process_next
from services.storage import database, document_dict, initialize

RECEIPT = "ORNEK MARKET\nVERGİ DAİRESİ: Kadıköy\nVKN: 1234567890\nFIS NO: 000123\nTARIH: 13.09.2026 SAAT: 14:25:36\nDEFTER %20 *120,00\nKDV %20\nTOPKDV 20,00\nTOPLAM 120,00\nNAKIT 120,00"
Z_REPORT = "ORNEK MARKET\nVERGİ DAİRESİ: Kadıköy\nVKN: 1234567890\nZ RAPORU NO: 000456\nMALİ SİCİL NO: AB00000123\nFİŞ ADEDİ: 2\nTARIH: 13.09.2026 SAAT: 23:15:42\nKDV %20 100,00 20,00\nSATIS TOPLAMI 120,00\nNAKIT 50,00\nKREDI KARTI 70,00"


def profile(program="luca"):
    return validate_profile({"program": program})


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.delenv("ALLOW_LOCAL_AUTH", raising=False)
    monkeypatch.setattr(server, "DB_PATH", tmp_path / "test.db")
    monkeypatch.setattr(server, "UPLOAD_DIR", tmp_path / "uploads")
    server.UPLOAD_DIR.mkdir()
    initialize(server.DB_PATH)
    server.app.config.update(TESTING=True, WORKER_ENABLED=False)
    private = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    monkeypatch.setattr(server.jwks, "get_signing_key_from_jwt", lambda token: SimpleNamespace(key=private.public_key()))

    def headers(user="alice", **extra):
        claims = {"sub": user, "aud": server.PROJECT_ID, "iss": f"https://securetoken.google.com/{server.PROJECT_ID}", "iat": int(time.time()), "exp": int(time.time()) + 3600, **extra}
        return {"Authorization": "Bearer " + jwt.encode(claims, private, algorithm="RS256")}

    result = server.app.test_client()
    result.headers_for = headers
    result.auth = headers()
    assert result.put("/api/settings", headers=result.auth, json=profile()).status_code == 200
    return result


def image_bytes(index=0):
    stream = io.BytesIO()
    Image.new("RGB", (30, 30), (index % 256, index // 256, 100)).save(stream, "PNG")
    stream.seek(0)
    return stream


def upload(client, index=0, kind="receipts", name=None, source=None):
    response = client.post("/api/documents", headers=client.auth, data={"kind": kind, "file": (source or image_bytes(index), name or f"fis-{index}.png")})
    assert response.status_code == 202, response.json
    return response.json["items"][0]["id"]


def successful_read(path, page, kind, attempt):
    return {**extract_document(RECEIPT if kind == "receipts" else Z_REPORT, kind), "confidence": 96, "engine": "TEST FIXTURE"}


@pytest.mark.parametrize("value,expected", [("1.234,56", "1234.56"), ("1234.56", "1234.56"), (1234.56, "1234.56"), ("120,00", "120.00"), ("1,234.56", "1234.56"), ("12.345.678,90", "12345678.90"), ("0", "0.00")])
def test_money(value, expected):
    assert decimal_money(value) == Decimal(expected)


@pytest.mark.parametrize("value", [None, "", "Bulunamadı", "NaN", "-", True, "12abc"])
def test_missing_is_not_zero(value):
    with pytest.raises(ValueError):
        decimal_money(value)


def test_extraction_and_tax_reconciliation():
    data = extract_document(RECEIPT, "receipts")
    assert data["issues"] == []
    assert data["vat_breakdown"] == [{"rate":20, "base":"100.00", "tax":"20.00"}]
    assert extract_document(RECEIPT.replace("20,00\nTOPLAM", "25,00\nTOPLAM"), "receipts")["issues"]
    assert extract_document(RECEIPT.replace("13.09.2026", "31.02.2026"), "receipts")["issues"]
    assert extract_document(RECEIPT.replace("TOPLAM 120,00", "TOPLAM KDV 20,00"), "receipts")["total_amount"] == ""
    assert extract_document(RECEIPT + "\nODENECEK 125,00", "receipts")["issues"]
    assert extract_document(RECEIPT + "\nFIS NO: 456", "receipts")["issues"]


def test_multi_vat_is_not_replaced_by_first_rate():
    text = RECEIPT.replace("KDV %20\nTOPKDV 20,00", "KDV %20 100,00 20,00\nKDV %10 100,00 10,00\nTOPKDV 30,00").replace("120,00", "230,00")
    text = text.replace("DEFTER %20 *230,00", "DEFTER %20 *120,00\nKİTAP %10 *110,00")
    data = extract_document(text, "receipts")
    assert not data["issues"]
    assert len(data["vat_breakdown"]) == 2
    assert extract_document(text.replace("KDV %10 100,00 10,00", "URUN %10"), "receipts")["issues"]


def test_z_payment_reconciliation():
    data = extract_document(Z_REPORT, "z-reports")
    assert not data["issues"], data["issues"]
    assert extract_document(Z_REPORT.replace("70,00", "60,00"), "z-reports")["issues"]
    assert extract_document(Z_REPORT, "receipts")["issues"]


def test_receipt_payment_conflict_and_multiple_dates_are_reviewed():
    assert extract_document(RECEIPT.replace("NAKIT 120,00", "NAKIT 100,00"), "receipts")["issues"]
    assert extract_document(RECEIPT + "\nTARIH: 14.09.2026", "receipts")["issues"]


def test_authentication_and_user_separation(client):
    doc = upload(client)
    assert client.get("/api/documents", headers={"X-Firebase-UID":"alice"}).status_code == 401
    assert client.get("/api/documents", headers=client.headers_for(exp=int(time.time()) - 1)).status_code == 401
    assert client.get("/api/documents", headers=client.headers_for(aud="wrong-project")).status_code == 401
    bob = client.headers_for("bob")
    assert client.get("/api/documents", headers=bob).json["items"] == []
    for path in (f"/api/documents/{doc}", f"/api/documents/{doc}/source"):
        assert client.get(path, headers=bob).status_code == 404
    assert client.post(f"/api/documents/{doc}/retry", headers=bob).status_code == 404
    assert client.get("/api/settings", headers=bob).json["selected"] is None


def test_125_uploads_sequential_queue_and_pagination(client):
    for index in range(125):
        upload(client, index)
    data = client.get("/api/documents", headers=client.auth).json
    assert data["total"] == 125 and data["counts"]["queued"] == 125 and len(data["items"]) == 50
    assert len(client.get("/api/documents?offset=100", headers=client.auth).json["items"]) == 25
    counter = 0

    def reader(*args):
        nonlocal counter
        counter += 1
        data = successful_read(*args)
        data["document_no"] = str(counter)
        return data

    while process_next(server.DB_PATH, reader):
        pass
    assert counter == 125
    assert client.get("/api/documents", headers=client.auth).json["counts"] == {"success":125}
    output = client.get("/export/excel?type=receipts", headers=client.auth)
    assert output.status_code == 200
    book = load_workbook(io.BytesIO(output.data))
    assert book.active.max_row == 377
    assert sum(row[7] or 0 for row in list(book.active.values)[1:-1]) == sum(row[8] or 0 for row in list(book.active.values)[1:-1]) == 15000


def test_failed_retry_survives_reload_and_deduplicates(client):
    doc = upload(client)
    assert upload(client) == doc

    def fails(*args):
        raise RuntimeError("Örnek okuma hatası")

    process_next(server.DB_PATH, fails)
    assert client.get(f"/api/documents/{doc}", headers=client.auth).json["status"] == "failed"
    assert client.post(f"/api/documents/{doc}/retry", headers=client.auth).status_code == 200
    assert client.post(f"/api/documents/{doc}/retry", headers=client.auth).status_code == 409
    initialize(server.DB_PATH)
    process_next(server.DB_PATH, successful_read)
    assert client.get(f"/api/documents/{doc}", headers=client.auth).json["attempts"] == 2
    second = upload(client, 1)
    process_next(server.DB_PATH, successful_read)
    data = client.get(f"/api/documents/{second}", headers=client.auth).json
    assert data["status"] == "duplicate" and data["duplicate_of"] == doc


def test_expired_worker_lease_is_recovered(client):
    doc = upload(client)
    with database(server.DB_PATH) as db:
        db.execute("UPDATE documents SET status='processing',started_at=? WHERE id=?", (time.time() - 1000, doc))
    process_next(server.DB_PATH, successful_read)
    assert client.get(f"/api/documents/{doc}", headers=client.auth).json["status"] == "success"


def test_export_blocks_pending_and_excludes_review_and_other_kind(client):
    upload(client)
    assert client.get("/export/excel?type=receipts", headers=client.auth).status_code == 409
    process_next(server.DB_PATH, successful_read)
    reviewed = upload(client, 1)
    process_next(server.DB_PATH, lambda *args: extract_document("unreadable", "receipts"))
    upload(client, 2, "z-reports")
    assert client.get("/export/excel?type=receipts", headers=client.auth).status_code == 200
    assert client.get("/export/excel?type=z-reports", headers=client.auth).status_code == 409
    process_next(server.DB_PATH, successful_read)
    book = load_workbook(io.BytesIO(client.get("/export/excel?type=receipts", headers=client.auth).data))
    assert book.active.max_row == 5
    assert client.get("/export/excel?type=all", headers=client.auth).status_code == 400


def test_template_maps_columns_preserves_ids_and_prevents_formula_execution(client):
    upload(client)
    process_next(server.DB_PATH, successful_read)
    changed = profile("eta")
    changed["template_name"] = "eta-test.xlsx"
    changed["header_row"] = 3
    changed["columns"] = list(reversed(changed["columns"]))
    assert client.put("/api/settings", headers=client.auth, json=changed).status_code == 200
    response = client.get("/export/excel?type=receipts&mapped=1", headers=client.auth)
    book = load_workbook(io.BytesIO(response.data))
    assert "Belge Bilgileri" in book.sheetnames and book.active.cell(3,1).value == "Kaynak Dosya"
    assert book.active.cell(4,9).value == "000123" and book.active.cell(4,9).data_type == "s"
    template = client.post("/api/template", headers=client.auth, data={"file":(io.BytesIO(response.data),"sample.xlsx")})
    assert template.status_code == 200 and template.json["header_row"] == 3
    document = {"kind":"receipts", "direction":"expense", "filename":"=HYPERLINK(bad).png", "result":successful_read(None,0,"receipts",1)}
    document["result"]["items"][0]["name"] = '=1+1'
    book = load_workbook(export_workbook([document], profile()))
    assert book.active["G2"].value == '=1+1' and book.active["G2"].data_type == 's'
    assert book.active["F2"].data_type == 's'


def test_income_and_expense_have_different_payment_accounts():
    configured = profile()
    data = extract_document(RECEIPT.replace("NAKIT 120,00", "KREDI KARTI 120,00"), "receipts")
    doc = {"kind":"receipts", "direction":"expense", "filename":"a.png", "result":data}
    expense = journal_rows([doc], configured)
    doc["direction"] = "income"
    income = journal_rows([doc], configured)
    assert expense[-1]["account"] == configured["accounts"]["expense_card"]
    assert income[-1]["account"] == configured["accounts"]["income_card"]
    assert expense[-1]["credit"] == income[-1]["debit"] == Decimal("120.00")


def test_pdf_upload_creates_one_job_per_page(client):
    pdf = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(pdf, "PDF", save_all=True, append_images=[Image.new("RGB", (100,100), "gray")])
    pdf.seek(0)
    upload(client, name="iki-sayfa.pdf", source=pdf)
    docs = client.get("/api/documents", headers=client.auth).json["items"]
    assert len(docs) == 2 and {doc["page"] for doc in docs} == {0, 1}


def test_rejects_invalid_file_and_requires_setup(client):
    bad = client.post("/api/documents", headers=client.auth, data={"file":(io.BytesIO(b'<script>bad</script>'),"fake.png")})
    assert bad.status_code == 400
    assert not list(server.UPLOAD_DIR.iterdir())
    assert client.post("/api/documents", headers=client.headers_for("bob"), data={"file":(image_bytes(),"x.png")}).status_code == 400
    assert client.put("/api/settings", headers=client.auth, json={"program":"unknown"}).status_code == 400
    assert client.get("/").status_code == 200
    assert client.get("/app").status_code == 200


@pytest.mark.parametrize("program", [program["id"] for program in PROGRAMS])
def test_program_selection_alone_loads_codes_and_exports_them(client, program):
    response = client.put("/api/settings", headers=client.auth, json={"program":program})
    assert response.status_code == 200
    assert response.json["profile"]["accounts"] == STANDARD_ACCOUNTS
    assert response.json["profile"]["account_plan"]["id"] == "tdhp-main-v1"
    upload(client)
    process_next(server.DB_PATH, successful_read)
    output = client.get("/export/excel?type=receipts", headers=client.auth)
    assert output.status_code == 200
    book = load_workbook(io.BytesIO(output.data))
    assert [book.active.cell(row,1).value for row in (2,3,4)] == ["770","191","100"]
    assert list(book["Aktarım Bilgisi"].values)[0][1] == next(p["name"] for p in PROGRAMS if p["id"] == program)


def test_old_manual_codes_cannot_override_automatic_codes(client):
    old = profile()
    old["accounts"] = {key:"999.OLD" for key in STANDARD_ACCOUNTS}
    old["version"] = "Existing version"
    old["template_name"] = "existing.xlsx"
    with database(server.DB_PATH) as db:
        db.execute("UPDATE profiles SET data=? WHERE user_id='alice' AND program='luca'", (json.dumps(old),))
    settings = client.get("/api/settings", headers=client.auth).json
    for value in (settings["selected"], settings["profiles"]["luca"]):
        assert value["accounts"] == STANDARD_ACCOUNTS
        assert value["version"] == "Existing version" and value["template_name"] == "existing.xlsx"
    doc_id = upload(client)
    process_next(server.DB_PATH, successful_read)
    with database(server.DB_PATH) as db:
        db.execute("UPDATE documents SET account_code='999.ROW' WHERE id=?", (doc_id,))
    output = client.get("/export/excel?type=receipts&mapped=1", headers=client.auth)
    assert output.status_code == 200
    book = load_workbook(io.BytesIO(output.data))
    assert [book.active.cell(row,4).value for row in (2,3,4)] == ["770","191","100"]
    response = client.put("/api/settings", headers=client.auth, json=old)
    assert response.json["profile"]["accounts"] == STANDARD_ACCOUNTS


def test_program_defaults_are_independent_copies():
    one = default_profile("luca")
    one["accounts"]["expense"] = "999"
    one["columns"][0]["header"] = "Changed"
    one["account_plan"]["name"] = "Changed"
    two = default_profile("zirve")
    assert two["accounts"]["expense"] == "770"
    assert two["columns"][0]["header"] == "Tarih"
    assert two["account_plan"]["name"] == "Tek Düzen · standart ana hesaplar"


def test_ocr_disagreement_never_gets_tick(tmp_path, monkeypatch):
    from services import document_ocr
    path = tmp_path / "fixture.png"
    path.write_bytes(image_bytes().read())
    monkeypatch.setattr(document_ocr.pytesseract, "get_languages", lambda **kwargs: ["eng", "tur"])
    count = 0

    def tokens(*args, **kwargs):
        nonlocal count
        count += 1
        text = RECEIPT if count == 1 else RECEIPT.replace("000123", "000124")
        words, line_numbers = [], []
        for number, line in enumerate(text.splitlines()):
            for word in line.split():
                words.append(word)
                line_numbers.append(number)
        return {"text":words, "conf":[98]*len(words), "block_num":[0]*len(words), "par_num":[0]*len(words), "line_num":line_numbers}

    monkeypatch.setattr(document_ocr.pytesseract, "image_to_data", tokens)
    assert read_document(path, 0, "receipts", 1)["issues"]


@pytest.mark.parametrize("lines,expected", [
    ("ÇİZGİLİ DEFTER %20 *120,00", "ÇİZGİLİ DEFTER"),
    ("ÇİZGİLİ DEFTER\n2 ADET X 60,00\n%20 *120,00", "ÇİZGİLİ DEFTER"),
    ("KALEM 2 X 60,00 %20 *120,00", "KALEM"),
    ("SU 1,5 L %20 *120,00", "SU 1,5 L"),
    ("ÇİZGİLİ DEFTER %20 *60,00\nMAVİ KALEM %20 *60,00", "ÇİZGİLİ DEFTER; MAVİ KALEM"),
    ("KALEM %20 *60,00\nKALEM %20 *60,00", "KALEM"),
])
def test_product_names_are_separate_from_seller_and_preserve_turkish(lines, expected):
    data = extract_document(RECEIPT.replace("DEFTER %20 *120,00", lines), "receipts")
    assert not data["issues"], data["issues"]
    assert data["seller_name"] == "ORNEK MARKET"
    assert data["product_name"] == expected
    document = {"kind": "receipts", "direction": "expense", "filename": "fis.png", "result": data}
    assert {row["description"] for row in journal_rows([document], profile())} == {expected}


@pytest.mark.parametrize("lines", [
    "", "%20 *120,00", "DEFTER %20 *100,00",
    "DEFTER %20 *140,00\nİNDİRİM -10,00",
    "ATATÜRK MAH. NO: 12\n%20 *120,00",
    "DEFTER %20 *80,00\nKALEM %20 *40,00\nSİLGİ %20 *10,00",
])
def test_incomplete_or_ambiguous_product_lines_need_review(lines):
    data = extract_document(RECEIPT.replace("DEFTER %20 *120,00", lines), "receipts")
    assert data["issues"]
    assert "ORNEK MARKET" not in data["product_name"]


def test_product_rate_totals_are_checked():
    text = RECEIPT.replace("DEFTER %20 *120,00", "KALEM %20 *110,00\nKİTAP %10 *120,00")
    text = text.replace("KDV %20\nTOPKDV 20,00", "KDV %20 100,00 20,00\nKDV %10 100,00 10,00\nTOPKDV 30,00")
    text = text.replace("TOPLAM 120,00", "TOPLAM 230,00").replace("NAKIT 120,00", "NAKIT 230,00")
    assert any("oranlarına" in issue for issue in extract_document(text, "receipts")["issues"])


@pytest.mark.parametrize("label", ["SERİ", "B. SERİ", "BELGE SERİSİ"])
def test_document_series_does_not_become_a_product_name(label):
    data = extract_document(RECEIPT.replace("DEFTER", f"{label}: AB\nDEFTER"), "receipts")
    assert not data["issues"]
    assert data["document_series"] == "AB" and data["product_name"] == "DEFTER"


def test_receipt_export_matches_supplied_column_order_and_totals():
    data = extract_document(RECEIPT.replace("DEFTER", "ÇİZGİLİ DEFTER"), "receipts")
    doc = {"kind": "receipts", "direction": "expense", "filename": "fis.png", "result": data}
    book = load_workbook(export_workbook([doc], profile()))
    sheet = book["Fiş"]
    assert list(next(sheet.values)) == ["Hesap Kodu", "Hesap Adı", "Belge Tipi", "Belge Tarihi", "B. Seri", "Belge No", "Açıklama", "Borç Tutar", "Alacak Tutar", "Hesap Adı", "."]
    assert sheet["B2"].value == sheet["J2"].value == "Genel Yönetim Giderleri"
    assert sheet["C2"].value == "Yazar Kasa Fişi"
    assert sheet["D2"].value.strftime("%d.%m.%Y") == "13.09.2026"
    assert sheet["E2"].value is None and sheet["F2"].value == "000123"
    assert all(sheet.cell(row, 7).value == "ÇİZGİLİ DEFTER" for row in (2, 3, 4))
    assert sheet["A5"].value == 3 and sheet["H5"].value == sheet["I5"].value == 120
    assert sheet.auto_filter.ref == "A1:K4"


def test_z_export_has_own_columns_and_balanced_payments():
    data = extract_document(Z_REPORT, "z-reports")
    doc = {"kind": "z-reports", "direction": "income", "filename": "z.png", "result": data}
    book = load_workbook(export_workbook([doc], profile()))
    sheet = book["Z Raporu"]
    assert "Fiş" not in book.sheetnames
    assert list(next(sheet.values)) == ["Hesap Kodu", "Hesap Adı", "Belge Tipi", "Belge Tarihi", "Rapor No", "Açıklama", "Borç Tutar", "Alacak Tutar", "."]
    assert [sheet.cell(row, 1).value for row in range(2, 6)] == ["600", "391", "100", "108"]
    assert sheet["E2"].value == "000456" and sheet["F2"].value == "Z Raporu 000456"
    assert sheet["G4"].value == 50 and sheet["G5"].value == 70
    assert sheet["G6"].value == sheet["H6"].value == 120


def test_template_recognizes_real_headers_after_annotations_and_duplicate_names():
    source = Workbook()
    sheet = source.active
    sheet.title = "Fiş"
    sheet.append(["Bu örnek bir açıklamadır"] * 11)
    headers = ["Hesap Kodu", "Hesap Adı", "Belge Tipi", "Belge Tarihi", "B. Seri", "Belge No", "Açıklama", "Borç Tutar", "Alacak Tutar", "Hesap Adı", "."]
    sheet.append(headers)
    output = io.BytesIO()
    source.save(output)
    output.seek(0)
    template = read_template(output)
    assert template["header_row"] == 2 and template["template_kind"] == "receipts"
    assert [col["field"] for col in template["columns"]] == ["account", "account_name", "document_type", "date", "document_series", "document_no", "description", "debit", "credit", "account_name", "blank"]


def test_receipt_template_cannot_leak_into_z_export(client):
    changed = profile()
    changed["template_name"] = "fis.xlsx"
    client.put("/api/settings", headers=client.auth, json=changed)
    upload(client, kind="z-reports")
    process_next(server.DB_PATH, successful_read)
    assert client.get("/export/excel?type=z-reports&mapped=1", headers=client.auth).status_code == 400
    output = client.get("/export/excel?type=z-reports", headers=client.auth)
    assert output.status_code == 200
    assert load_workbook(io.BytesIO(output.data)).active.title == "Z Raporu"


def test_legacy_seller_descriptions_require_retry_without_losing_source(client):
    doc_id = upload(client)
    process_next(server.DB_PATH, successful_read)
    with database(server.DB_PATH) as db:
        old = json.loads(db.execute("SELECT result FROM documents WHERE id=?", (doc_id,)).fetchone()[0])
        old.pop("extraction_version")
        old.pop("seller_name")
        old.pop("items")
        old["product_name"] = "ORNEK MARKET"
        db.execute("UPDATE documents SET result=? WHERE id=?", (json.dumps(old), doc_id))
    initialize(server.DB_PATH)
    initialize(server.DB_PATH)  # Migration is idempotent.
    doc = client.get(f"/api/documents/{doc_id}", headers=client.auth).json
    assert doc["status"] == "review" and doc["result"]["product_name"] == ""
    assert doc["result"]["seller_name"] == "ORNEK MARKET"
    assert doc["result"]["raw_text"] == RECEIPT
    assert client.get("/export/excel", headers=client.auth).status_code == 400
    with pytest.raises(ValueError, match="Ürün adları"):
        export_workbook([{"kind":"receipts", "direction":"expense", "filename":"old.png", "result":old}], profile())
    assert client.post(f"/api/documents/{doc_id}/retry", headers=client.auth).status_code == 200
    process_next(server.DB_PATH, successful_read)
    assert client.get(f"/api/documents/{doc_id}", headers=client.auth).json["status"] == "success"


def test_ocr_product_name_disagreement_is_reviewed(tmp_path, monkeypatch):
    from services import document_ocr
    path = tmp_path / "fixture.png"
    path.write_bytes(image_bytes().read())
    monkeypatch.setattr(document_ocr.pytesseract, "get_languages", lambda **kwargs: ["eng", "tur"])
    texts = iter([RECEIPT, RECEIPT.replace("DEFTER", "KALEM")])
    def tokens(*args, **kwargs):
        words, numbers = [], []
        for index, line in enumerate(next(texts).splitlines()):
            words.extend(line.split())
            numbers.extend([index] * len(line.split()))
        return {"text":words, "conf":[98]*len(words), "block_num":[0]*len(words), "par_num":[0]*len(words), "line_num":numbers}
    monkeypatch.setattr(document_ocr.pytesseract, "image_to_data", tokens)
    result = read_document(path, 0, "receipts", 1)
    assert any("İki okuma" in issue for issue in result["issues"])

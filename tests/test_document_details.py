import io
import json
from datetime import datetime, time
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from test_workflow import RECEIPT, Z_REPORT, client, image_bytes, profile, server, successful_read, upload
from test_banking import configured_chart
from services.accounting import export_workbook, journal_rows
from services.document_extraction import extract_document
from services.document_ocr import read_tesseract_document as read_document
from services.queue_worker import process_next
from services.storage import database, initialize


def doc(text=RECEIPT, kind="receipts"):
    return {"kind": kind, "direction": "income" if kind == "z-reports" else "expense", "filename": "TEST-belge.png", "result": extract_document(text, kind)}


def sheet_rows(book, title):
    rows = list(book[title].values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


@pytest.mark.parametrize("label", ["VERGİ DAİRESİ: Kadıköy", "Kadıköy V.D. 1234567890", "Kadıköy V D 1234567890", "VERGİ DAİRESİ Kadıköy", "V.D.: Kadıköy"])
def test_tax_office_preserves_turkish_name(label):
    data = extract_document(RECEIPT.replace("VERGİ DAİRESİ: Kadıköy", label), "receipts")
    assert not data["issues"] and data["tax_office"] == "Kadıköy"
    assert data["document_time"] == "14:25:36" and data["document_datetime"] == "2026-09-13T14:25:36"


@pytest.mark.parametrize("label,expected", [
    ("$15L1 V D 1234567890", "SISLI"),
    ("$1$L1 V.D. 1234567890", "SISLI"),
    ("VERGİ DAİRESİ: KAD1KOY", "KADIKOY"),
    ("VERGİ DAİRESİ: ŞİŞLİ", "ŞİŞLİ"),
    ("VERGİ DAİRESİ: 19 MAYIS", "19 MAYIS"),
    ("VERGİ DAİRESİ: 19MAYIS", "19MAYIS"),
    ("VERGİ DAİRESİ: 15 TEMMUZ", "15 TEMMUZ"),
])
def test_office_glyph_repair_is_scoped_to_explicit_name(label, expected):
    text = RECEIPT.replace("VERGİ DAİRESİ: Kadıköy", label)
    data = extract_document(text, "receipts")
    assert not data['issues'], data['issues']
    assert data['tax_office'] == expected
    assert data['tax_id'] == '1234567890'
    assert data['document_datetime'] == '2026-09-13T14:25:36'
    assert data['total_amount'] == '120.00'
    assert data['raw_text'] == text
    if '$' in label or 'KAD1KOY' in label:
        assert any('OCR karakterleri düzeltildi' in note for note in data['notes'])


def test_office_glyph_repair_does_not_guess_from_address_or_repair_tax_id():
    text = RECEIPT.replace('VERGİ DAİRESİ: Kadıköy', 'CADDE NO 1 $15L1/ISTANBUL')
    assert extract_document(text, 'receipts')['tax_office'] == ''
    text = RECEIPT.replace('VERGİ DAİRESİ: Kadıköy', '$15L1 V D 123456789O').replace('1234567890', '123456789O')
    assert extract_document(text, 'receipts')['tax_id'] == ''


@pytest.mark.parametrize("source,part,label", [(RECEIPT, "VERGİ DAİRESİ: Kadıköy\n", "Vergi dairesi"), (RECEIPT, " SAAT: 14:25:36", "Saat"),
                                              (Z_REPORT, "MALİ SİCİL NO: AB00000123\n", "sicil"), (Z_REPORT, "FİŞ ADEDİ: 2\n", "adedi")])
def test_missing_required_details_block_tick_and_export(source, part, label):
    document = doc(source.replace(part, ""), "z-reports" if source == Z_REPORT else "receipts")
    assert any(label in issue for issue in document["result"]["issues"])
    with pytest.raises(ValueError):
        export_workbook([document], profile())
    if label == "Saat":
        assert document["result"]["document_datetime"] == "2026-09-13"


@pytest.mark.parametrize("extra", ["\nSAAT: 14:25:37", "\nVERGİ DAİRESİ: Üsküdar"])
def test_conflicting_metadata_is_reviewed(extra):
    assert any("çelişiyor" in issue for issue in doc(RECEIPT + extra)["result"]["issues"])


def test_labeled_multiline_business_title_stays_separate_from_products():
    data = doc(RECEIPT.replace("ORNEK MARKET", "İŞLETME UNVANI: ÖRNEK TİCARET\nLİMİTED ŞİRKETİ"))["result"]
    assert not data["issues"]
    assert data["seller_name"] == "ÖRNEK TİCARET LİMİTED ŞİRKETİ" and data["product_name"] == "DEFTER"


@pytest.mark.parametrize("product,quantity,unit,price", [
    ("ÇİZGİLİ DEFTER 2 ADET X 60,00 %20 *120,00", "2", "adet", "60.00"),
    ("ÇİZGİLİ DEFTER\n2 ADET X 60,00\n%20 *120,00", "2", "adet", "60.00"),
    ("ELMA 1,500 KG X 80,00 %20 *120,00", "1.500", "kg", "80.00"),
])
def test_quantity_unit_price_and_product_amount_are_preserved(product, quantity, unit, price):
    document = doc(RECEIPT.replace("DEFTER %20 *120,00", product))
    data = document["result"]
    assert not data["issues"], data["issues"]
    item = data["items"][0]
    assert (item["quantity"], item["unit"], item["unit_price"], item["amount"]) == (quantity, unit, price, "120.00")
    row = sheet_rows(load_workbook(export_workbook([document], profile())), "Fiş Kalemleri")[0]
    assert row["Miktar"] == float(quantity) and row["Birim fiyat"] == float(price) and row["KDV oranı (%)"] == 20


def test_absent_optional_quantity_is_not_fabricated_and_rate_is_not_quantity():
    data = doc()["result"]
    assert not data["issues"]
    assert data["items"][0]["quantity"] == data["items"][0]["unit_price"] == ""
    row = sheet_rows(load_workbook(export_workbook([doc()], profile())), "Fiş Kalemleri")[0]
    assert row["Miktar"] is None and row["Birim fiyat"] is None


def test_wrong_multiplication_and_missing_item_vat_are_reviewed():
    assert doc(RECEIPT.replace("DEFTER %20", "DEFTER 2 X 70,00 %20"))["result"]["issues"]
    assert doc(RECEIPT.replace("DEFTER %20", "DEFTER"))["result"]["issues"]
    assert doc(RECEIPT.replace("DEFTER %20", "DEFTER 2 X 60,0000 %20"))["result"]["issues"]


def test_iso_dates_are_checked_for_conflicts_without_rejecting_equivalent_dates():
    assert any("birden fazla tarih" in issue for issue in doc(RECEIPT + "\nTARİH: 2026-09-14")["result"]["issues"])
    assert not doc(RECEIPT + "\nTARİH: 2026-09-13")["result"]["issues"]


def test_starred_money_is_not_a_bank_card_identifier():
    from services.banking import bank_evidence
    assert bank_evidence("TOPLAM ****1200,00\nKART NO ****1234")["card_last4"] == ["1234"]


def test_single_rate_discount_reconciles_without_double_subtraction():
    text = RECEIPT.replace("DEFTER %20 *120,00", "ÇİZGİLİ DEFTER 2 ADET X 60,00 %20 *120,00\nİNDİRİM -12,00")
    text = text.replace("TOPKDV 20,00", "TOPKDV 18,00").replace("TOPLAM 120,00", "TOPLAM 108,00").replace("NAKIT 120,00", "NAKIT 108,00")
    document = doc(text)
    assert not document["result"]["issues"], document["result"]["issues"]
    assert document["result"]["adjustments"]["discount"] == {"amount": "12.00", "count": None}
    entries = journal_rows([document], profile())
    assert entries[0]["debit"] == Decimal("90.00") and entries[-1]["credit"] == Decimal("108.00")
    assert all(row["description"] == "ÇİZGİLİ DEFTER" for row in entries)


def rich_z():
    return Z_REPORT.replace("KREDI KARTI 70,00", "KREDI KARTI 20,00\nYEMEK KARTI 50,00\nMULTINET 50,00") + "\nCIHAZ NO: 00001234\nIPTAL 2 24,00\nIADE ADET: 1 TUTAR: 12,00\nINDIRIM ADET: 0 TUTAR: 0,00\nKUMULATIF SATIS 12.000,00\nKUMULATIF KDV 2.000,00"


def test_z_details_meal_payments_and_all_excel_fields():
    document = doc(rich_z(), "z-reports")
    data = document["result"]
    assert not data["issues"], data["issues"]
    assert data["fiscal_id"] == "AB00000123" and data["device_no"] == "00001234"
    assert data["transaction_count"] == 2 and data["meal_card_amount"] == "50.00"
    assert len(data["payment_entries"]) == 3
    assert data["adjustments"]["cancellation"] == {"amount": "24.00", "count": 2}
    book = load_workbook(export_workbook([document], profile()))
    assert book.active.title == "Z Raporu" and "Fiş Kalemleri" not in book.sheetnames
    info = sheet_rows(book, "Belge Bilgileri")[0]
    assert info["VKN / TCKN"] == "1234567890" and info["Z raporu numarası"] == "000456"
    assert info["Saat"] == time(23, 15, 42) and info["Tarih"] == datetime(2026, 9, 13)
    assert info["Cihaz numarası"] == "00001234" and info["Kümülatif satış"] == 12000 and info["Kümülatif KDV"] == 2000
    assert info["İptal adedi"] == 2 and info["İade adedi"] == 1 and info["İndirim tutarı"] == 0
    vat = sheet_rows(book, "KDV Dağılımı")[0]
    assert [vat[key] for key in ("KDV dâhil satış tutarı", "KDV hariç matrah", "KDV tutarı")] == [120, 100, 20]
    payments = sheet_rows(book, "Ödeme Dağılımı")
    assert [row["Ödeme şekli"] for row in payments] == ["Nakit", "Kredi kartı", "Yemek kartı"]
    assert payments[-1]["Tutar"] == 50 and payments[-1]["Yemek kartı kuruluşu"] == "MULTINET"
    assert len(sheet_rows(book, "İndirim İptal İade")) == 3


def test_meal_card_uses_firm_named_account_not_bank_pos_account():
    configured = configured_chart()
    configured["chart"]["accounts"].append({"code": "120.001", "name": "Multinet yemek kartı", "role": "income_other", "leaf": True, "currency": "TRY", "bank_code": "", "card_last4": "", "iban": "", "rate": None, "categories": []})
    document = doc(Z_REPORT.replace("KREDI KARTI 70,00", "MULTINET 70,00"), "z-reports")
    assert not document["result"]["issues"]
    assert journal_rows([document], configured)[-1]["account"] == "120.001"


@pytest.mark.parametrize("change", [lambda text: text.replace("12.000,00", "100,00"), lambda text: text.replace("2.000,00", "10,00"),
                                    lambda text: text.replace("MULTINET 50,00", "MULTINET 40,00"), lambda text: text.replace("FİŞ ADEDİ: 2", "FİŞ ADEDİ: 0"),
                                    lambda text: text.replace("IPTAL 2 24,00", "IPTAL TUTARI OKUNAMIYOR")])
def test_conflicting_z_optional_fields_and_totals_block_readiness(change):
    assert doc(change(rich_z()), "z-reports")["result"]["issues"]


def test_z_three_column_vat_gross_base_and_tax_are_validated():
    data = doc(Z_REPORT.replace("KDV %20 100,00 20,00", "KDV %20 120,00 100,00 20,00"), "z-reports")["result"]
    assert not data["issues"]
    assert doc(Z_REPORT.replace("KDV %20 100,00 20,00", "KDV %20 125,00 100,00 20,00"), "z-reports")["result"]["issues"]


def test_v3_records_are_requeued_once_without_losing_source_or_history(client):
    identity = upload(client)
    process_next(server.DB_PATH, successful_read)
    with database(server.DB_PATH) as db:
        row = db.execute("SELECT * FROM documents WHERE id=?", (identity,)).fetchone()
        old = json.loads(row["result"])
        old["extraction_version"] = 3
        db.execute("UPDATE documents SET result=?, export_count=1 WHERE id=?", (json.dumps(old), identity))
        path = row["path"]
    initialize(server.DB_PATH)
    initialize(server.DB_PATH)
    with database(server.DB_PATH) as db:
        row = db.execute("SELECT * FROM documents WHERE id=?", (identity,)).fetchone()
        assert row["status"] == "queued" and row["path"] == path and row["export_count"] == 1
    process_next(server.DB_PATH, successful_read)
    initialize(server.DB_PATH)
    assert client.get(f"/api/documents/{identity}", headers=client.auth).json["status"] == "success"


def test_two_registers_with_same_z_number_are_not_duplicates(client):
    first = upload(client, kind="z-reports")
    process_next(server.DB_PATH, successful_read)
    second = upload(client, index=1, kind="z-reports")
    process_next(server.DB_PATH, lambda *args: extract_document(Z_REPORT.replace("AB00000123", "AB00000124"), "z-reports"))
    assert client.get(f"/api/documents/{second}", headers=client.auth).json["status"] == "success"


@pytest.mark.parametrize("replacement", [("Kadıköy", "Üsküdar"), ("AB00000123", "AB00000124"), ("23:15:42", "23:15:43"),
                                         ("FİŞ ADEDİ: 2", "FİŞ ADEDİ: 3"), ("2.000,00", "2.001,00"), ("IPTAL 2 24,00", "IPTAL 2 25,00")])
def test_new_details_must_agree_between_both_ocr_reads(tmp_path, monkeypatch, replacement):
    from services import document_ocr
    assert not extract_document(rich_z(), "z-reports")["issues"]
    path = tmp_path / "test.png"
    path.write_bytes(image_bytes().read())
    monkeypatch.setattr(document_ocr.pytesseract, "get_languages", lambda **kwargs: ["eng", "tur"])
    sources = iter([rich_z(), rich_z().replace(*replacement)])
    def tokens(*args, **kwargs):
        words, numbers = [], []
        for index, line in enumerate(next(sources).splitlines()):
            words.extend(line.split())
            numbers.extend([index] * len(line.split()))
        return {"text":words, "conf":[98]*len(words), "block_num":[0]*len(words), "par_num":[0]*len(words), "line_num":numbers}
    monkeypatch.setattr(document_ocr.pytesseract, "image_to_data", tokens)
    assert any("İki okuma" in issue for issue in read_document(path, 0, "z-reports", 1)["issues"])


def test_bracket_corrupted_topkdv_in_clean_line_and_extraction():
    from services.document_ocr import _clean_ocr_line
    assert _clean_ocr_line("[OPADV *758,33") == "TOPKDV *758,33"
    assert _clean_ocr_line("[OPKDV 100,00") == "TOPKDV 100,00"
    
    z_text = (
        "ORNEK MARKET\n"
        "VERGİ DAİRESİ: Kadıköy\n"
        "VKN: 1234567890\n"
        "Z RAPORU NO: 000456\n"
        "MALİ SİCİL NO: AB00000123\n"
        "FİŞ ADEDİ: 2\n"
        "TARIH: 13.09.2026 SAAT: 23:15:42\n"
        "GUNLUK FIS DOKUMU\n"
        "TOPLAM *4.550,00\n"
        "[OPADV *758,33\n"
        "KDV %20\n"
        "KREDI *4.550,00\n"
    )
    data = extract_document(z_text, "z-reports")
    assert not data["issues"], data["issues"]
    assert data["total_amount"] == "4550.00"
    assert data["vat_breakdown"] == [{"rate": 20, "base": "3791.67", "tax": "758.33"}]
    assert data["card_amount"] == "4550.00"


def test_z_report_counters_and_unreadable_datetime_are_not_fabricated():
    from services.document_ocr import _clean_ocr_line
    assert _clean_ocr_line("iakiv 0/i851/026") == "iakiv 0/i851/026"
    assert _clean_ocr_line("SAi 0181202") == "SAi 0181202"
    assert _clean_ocr_line("515LI V D 3880097945") == "515LI V D 3880097945"

    z_text = (
        "SAN. TIC. LTD. STI\n"
        "ERGENEKON MAH. CUMHURIYET CAD.\n"
        "FRANSIZ HASTANESI SK NO 349/1 SISLI / ISTANBUL\n"
        "SISLI V.D. 3880097945\n"
        "TARIH 07/05/2026\n"
        "SAAT 01:51:02\n"
        "Z RAPORU\n"
        "RAPOR NO 1759\n"
        "MALİ SİCİL NO: JH20011571\n"
        "FİŞ ADEDİ: 5\n"
        "MALI BELLEK TOPLAMI *9.615.138,88\n"
        "MALI BELLEK TOP KDV *1.558.529,41\n"
        "GUNLUK FIS DOKUMU\n"
        "TOPLAM *4.550,00\n"
        "TOPKDV *758,33\n"
        "KDV BILGILERI\n"
        "KDV %20.00 *758,33\n"
        "TOPLAM *4.550,00\n"
        "ODEME BILGILERI\n"
        "NAKIT *0,00\n"
        "KREDI 5\n"
        "TOPLAM *4.550,00\n"
        "SAYACLAR\n"
        "ARTTIRIM ADET 0\n"
        "ARTTIRIM TUTAR *0,00\n"
        "INDIRIM ADET 0\n"
        "INDIRIM TUTAR *0,00\n"
        "DUZELTME ADET 0\n"
        "DUZELTME TUTAR *0,00\n"
        "SATIS IPTAL 0\n"
        "SATIS IPTAL TUTAR *0,00\n"
        "EKU NO: 0001 Z NO: 1759\n"
    )
    data = extract_document(z_text, "z-reports")
    assert not data["issues"], data["issues"]
    assert data["total_amount"] == "4550.00"
    assert data["document_datetime"] == "2026-05-07T01:51:02"
    assert data["document_time"] == "01:51:02"
    assert data["tax_office"] == "SISLI"
    assert data["tax_id"] == "3880097945"
    assert data["card_amount"] == "4550.00"
    assert data["cash_amount"] == "0.00"
    assert not any("İndirim" in issue for issue in data["issues"])
    assert not any("Tarih" in issue for issue in data["issues"])
    assert not any("Saat" in issue for issue in data["issues"])


def test_z_report_bad_rates_and_conflicting_payment_sections_require_review():
    from services.document_ocr import _clean_ocr_line
    raw_ocr = (
        "FORA TURIZM REKLAM\n"
        "SAN. TİC LTD. STL.\n"
        "CUHURİYEL CAD. HO: 349/1\n"
        "SISLI/ İSTANBUL\n"
        "1 ŞİŞLİ 3880097945 ie\n"
        "TARIH - 30/05/2026\n"
        "SAAT - 04:21:42\n"
        "| FES NO > 0041\n"
        "Z RAPORU\n"
        "PAPOR No 1880\n"
        "PALL BELLEK TOPLAM = *12.867.454,44\n"
        "“MALİ BELLEK TOP KOV *2.070 030,56\n"
        "”--------GÜNLÜK FİŞ DÖKÜMÜ-------\n"
        "TOPLAM *35.650,00\n"
        "\"| Topkov *5.941,66\n"
        ": ----------KDV BİLGİLERİ --><-----\n"
        "* KOV 4270.00 Wan:\n"
        "TOPLAN “6 600\n"
        "-------DEPARTMAN BİLGİLERİ ------\n"
        "BIRA %20.00\n"
        "|, TOPLAM Ue *34.400,00\n"
        "| MIKTARI 32,0000\n"
        "YERLİ İÇKİ 420.00 z\n"
        "TOPLAN 41 250,00\n"
        "HIKTARI 1,0000\n"
        "--<------ÖDEHE BİLGİLERİ --------\n"
        "NAKIT 6\n"
        "TOPLAM *0,00\n"
        "KREDİ 33 5\n"
        "TOPLAM 50,0\n"
        "----------BELGE TİPLERİ---------\n"
        "OKC FİŞLERİ 33\n"
        "-KOV TOPLAMI *5 941,66 :\n"
        "-SATIŞ TOPLAMI *35.650,00\n"
        "-HAKİT *0,00\n"
        "“KREDİ *35.650,00\n"
        "“DİĞER *0 00\n"
        ". © - PRIPTAL 5\n"
        "ğ -KDV TOPLAMI *715,00\n"
        "Ş -SATIŞ TOPLAMI *4 650,00\n"
        "we ee------~=SAYACLAR---------->-\n"
        "ARTTIRIM ADET 0\n"
        "ARITIRIM TUTAR *0,00\n"
        "İNDİRİM ADET 0\n"
        "İNDİRİK TUTAR *0,00\n"
        "DÜZELTME ADET Q\n"
        ": DÜZELTME TUTAR *0 00\n"
        "HALİ FİŞ ADET 34\n"
        ": MÜŞTERİ FİŞİ ADETİ 33\n"
        "SE © HALİ OLMAYAN FİŞ 5\n"
        "SATIŞ İPTAL 5\n"
        "SATIS İPTAL TUTAR *4 650,00\n"
        "SERVIS 'HÜDAHALE 0\n"
        "....----- KASİYER BİLGİ---------\n"
        "KASIYERİ #35 650,00\n"
        "KASİYER: KASIYERİ :\n"
        "Tic. Sici]No: 344177\n"
        "EKÜ NO:0001 Z NO: 1880\n"
        "MF JH 20004135\n"
    )
    cleaned = "\n".join(_clean_ocr_line(l) for l in raw_ocr.splitlines())
    data = extract_document(cleaned, "z-reports")
    assert data["issues"]
    assert data["total_amount"] == "35650.00"
    assert data["document_no"] == "1880"
    assert data["tax_id"] == ""
    assert any("çelişiyor" in issue or "uyuşmuyor" in issue for issue in data["issues"])

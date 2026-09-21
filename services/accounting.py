from __future__ import annotations

import io
import re
import textwrap
from copy import deepcopy
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from services.document_extraction import EXTRACTION_VERSION, decimal_money, folded, receipt_description
from services.account_chart import select_account
from services.banking import BANK_BY_CODE, extract_payments

# Program selection loads a standard chart; vendor/company subaccounts are not universal.
PROGRAMS = [
    {"id": key, "name": name} for key, name in (
        ("luca", "Luca"), ("zirve", "Zirve"), ("logo", "Logo"), ("mikro", "Mikro"),
        ("eta", "ETA"), ("netsis", "Logo Netsis"), ("gms", "Mikrokom GMS.NET"),
        ("orka", "Orka"), ("datasoft", "Datasoft"), ("dia", "DİA"),
        ("zenom", "Zenom"), ("akinsoft", "Akınsoft"), ("custom", "Diğer / özel şablon")
    )
]
FIELDS = {
    "date": "Tarih", "voucher": "Fiş No", "document_no": "Belge No", "account": "Hesap Kodu",
    "description": "Açıklama", "debit": "Borç", "credit": "Alacak", "tax_id": "VKN / TCKN",
    "rate": "KDV Oranı", "direction": "Gelir / Gider", "filename": "Kaynak Dosya",
    "currency": "Para Birimi", "blank": "Boş bırak", "constant": "Sabit değer",
    "account_name": "Hesap Adı", "document_type": "Belge Tipi", "document_series": "B. Seri", "seller_name": "Firma Adı",
    "bank_code": "Banka EFT Kodu", "bank_name": "Banka Adı", "payment_method": "Ödeme Yöntemi", "card_last4": "Kart Son 4 Hane",
}
FIELDS.update({"tax_office": "Vergi Dairesi", "document_time": "Saat", "fiscal_id": "Mali Sicil Numarası", "device_no": "Cihaz Numarası",
               "transaction_count": "Fiş / İşlem Adedi", "cumulative_sales": "Kümülatif Satış", "cumulative_vat": "Kümülatif KDV",
               "total_amount": "Genel Toplam", "vat_amount": "Toplam KDV", "discount_amount": "İndirim Tutarı"})
DEFAULT_COLUMNS = [{"header": FIELDS[key], "field": key} for key in ("date", "voucher", "document_no", "account", "description", "debit", "credit", "tax_id", "rate", "direction", "filename")]
# Column order comes from the supplied Fiş / Z Raporu examples. Sample values,
# company-specific subaccounts, instructional rows and incorrect totals do not.
RECEIPT_COLUMNS = [{"header": header, "field": field} for header, field in (
    ("Hesap Kodu", "account"), ("Hesap Adı", "account_name"), ("Belge Tipi", "document_type"),
    ("Belge Tarihi", "date"), ("B. Seri", "document_series"), ("Belge No", "document_no"),
    ("Açıklama", "description"), ("Borç Tutar", "debit"), ("Alacak Tutar", "credit"),
    ("Hesap Adı", "account_name"), (".", "blank"),
)]
Z_COLUMNS = [{"header": header, "field": field} for header, field in (
    ("Hesap Kodu", "account"), ("Hesap Adı", "account_name"), ("Belge Tipi", "document_type"),
    ("Belge Tarihi", "date"), ("Rapor No", "document_no"), ("Açıklama", "description"),
    ("Borç Tutar", "debit"), ("Alacak Tutar", "credit"), (".", "blank"),
)]
ACCOUNT_LABELS = {"expense": "Gider hesabı", "income": "Gelir hesabı", "input_vat": "İndirilecek KDV hesabı", "output_vat": "Hesaplanan KDV hesabı", "cash": "Nakit hesabı", "expense_card": "Kartla yapılan ödeme hesabı", "income_card": "Kart / POS tahsilat hesabı", "expense_other": "Gider karşı hesabı (diğer / belirsiz ödeme)", "income_other": "Gelir karşı hesabı (diğer / belirsiz tahsilat)"}
ACCOUNT_LABELS["bank"] = "Banka kartı / havale / EFT hesabı"
STANDARD_ACCOUNTS = {
    "expense": "770", "income": "600", "input_vat": "191", "output_vat": "391",
    "cash": "100", "expense_card": "309", "income_card": "108",
    "expense_other": "320", "income_other": "120", "bank": "102",
}
ACCOUNT_NAMES = {
    "770": "Genel Yönetim Giderleri", "600": "Yurtiçi Satışlar", "191": "İndirilecek KDV",
    "391": "Hesaplanan KDV", "100": "Kasa", "309": "Diğer Mali Borçlar",
    "108": "Diğer Hazır Değerler", "320": "Satıcılar", "120": "Alıcılar",
    "102": "Bankalar",
}
ACCOUNT_PLAN = {
    "id": "tdhp-main-v1", "name": "Tek Düzen · standart ana hesaplar",
    "description": "Program seçiminizle otomatik yüklenir. Firmaya özel alt hesaplar içermez.",
    "source_url": "https://artvin.ticaret.gov.tr/data/642d4ed813b8768238acc3b1/c49f6df76c846dd1458ec9d4242aa3b3.pdf",
}


def resolve_profile(data):
    """Use the same server-owned codes for preview, saved settings and export.

    Older manually entered account settings stay in storage for traceability but
    cannot override automatic codes when a profile is read or used for export.
    """
    if not isinstance(data, dict) or data.get("program") not in {p["id"] for p in PROGRAMS}:
        raise ValueError("Muhasebe programını seçin.")
    profile = deepcopy(data)
    profile["accounts"] = dict(STANDARD_ACCOUNTS)
    profile["account_plan"] = dict(ACCOUNT_PLAN)
    profile["template_kind"] = data.get("template_kind", "receipts")
    profile["chart_id"] = data.get("chart_id", "")
    return profile


def default_profile(program):
    return resolve_profile({"program": program, "version": "", "columns": DEFAULT_COLUMNS,
                            "template_name": "", "sheet_name": "Muhasebe Fişleri", "header_row": 1, "date_format": "dd.mm.yyyy"})


def validate_profile(data):
    if not isinstance(data, dict) or data.get("program") not in {program["id"] for program in PROGRAMS}:
        raise ValueError("Muhasebe programını seçin.")
    profile = default_profile(data["program"])
    # Accounts are derived from the chosen profile, never from browser input.
    columns = data.get("columns", DEFAULT_COLUMNS)
    if not isinstance(columns, list) or not 6 <= len(columns) <= 80:
        raise ValueError("Şablon 6 ile 80 sütun içermeli.")
    for col in columns:
        if not isinstance(col, dict) or col.get("field") not in FIELDS or not str(col.get("header", "")).strip():
            raise ValueError("Excel sütun eşleştirmesi eksik.")
        if len(str(col.get("header"))) > 150 or len(str(col.get("value", ""))) > 150:
            raise ValueError("Excel sütun değeri çok uzun.")
    required = {"date", "document_no", "account", "debit", "credit"}
    if not required.issubset({col["field"] for col in columns}):
        raise ValueError("Tarih, belge no, hesap kodu, borç ve alacak sütunlarını eşleştirin.")
    profile["columns"] = columns
    profile["version"] = str(data.get("version", ""))[:100]
    profile["template_name"] = str(data.get("template_name", ""))[:200]
    profile["chart_id"] = str(data.get("chart_id", ""))
    if profile["chart_id"] and not re.fullmatch(r"[a-f0-9]{32}", profile["chart_id"]):
        raise ValueError("Firma hesap planı seçimi geçersiz.")
    profile["template_kind"] = data.get("template_kind", "receipts")
    if profile["template_kind"] not in {"receipts", "z-reports"}:
        raise ValueError("Şablon için fiş veya Z raporu seçin.")
    profile["sheet_name"] = re.sub(r"[\\/*?:\[\]]", "", str(data.get("sheet_name", "Muhasebe Fişleri")))[:31] or "Muhasebe Fişleri"
    profile["header_row"] = max(1, min(20, int(data.get("header_row", 1))))
    if data.get("date_format") in {"dd.mm.yyyy", "yyyy-mm-dd", "dd/mm/yyyy"}:
        profile["date_format"] = data["date_format"]
    return profile


def read_template(stream):
    import zipfile
    try:
        raw = stream.read(5 * 1024 * 1024 + 1)
        if len(raw) > 5 * 1024 * 1024:
            raise ValueError("Şablon en fazla 5 MB olabilir.")
        candidates = []
        if raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
            import xlrd
            with xlrd.open_workbook(file_contents=raw, on_demand=True) as book:
                for sheet in book.sheets():
                    candidates.extend((sheet.name, index + 1, sheet.row_values(index)[:80]) for index in range(min(20, sheet.nrows)))
        else:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("Şablon açıldığında çok büyük. Boş aktarım şablonunu yükleyin.")
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                for sheet in book.worksheets:
                    candidates.extend((sheet.title, index, list(row)) for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, max_col=80, values_only=True), 1))
            finally:
                book.close()
        aliases = {folded(label).replace(" ", ""): key for key, label in FIELDS.items()}
        aliases.update({"FISTARIHI": "date", "BELGETARIHI": "date", "BELGENUMARASI": "document_no", "RAPORNO": "document_no", "BORCTUTAR": "debit", "ALACAKTUTAR": "credit", "BORCTUTARI": "debit", "ALACAKTUTARI": "credit", "FISNUMARASI": "voucher"})
        def field(value):
            return aliases.get(folded(str(value or "")).replace(" ", ""), "blank")
        if not candidates:
            raise ValueError("Şablonda başlık bulunamadı.")
        candidates = [item for item in candidates if {"date", "document_no", "account", "debit", "credit"}.issubset({field(value) for value in item[2]})]
        if not candidates:
            raise ValueError("Şablonda tarih, belge/rapor no, hesap kodu, borç ve alacak başlıkları bulunamadı.")
        title, index, headers = max(candidates, key=lambda item: sum(field(value) != "blank" for value in item[2]))
        while headers and headers[-1] in (None, ""):
            headers.pop()
        if len(headers) < 6 or not {"date", "document_no", "account", "debit", "credit"}.issubset({field(value) for value in headers}):
            raise ValueError("Şablonda tarih, belge/rapor no, hesap kodu, borç ve alacak başlıkları bulunamadı.")
        columns = [{"header": str(header or f"Sütun {number}")[:150], "field": field(header)} for number, header in enumerate(headers, 1)]
        is_z = "Z RAPOR" in folded(title) or any(folded(str(value)) == "RAPOR NO" for value in headers)
        return {"columns": columns, "sheet_name": title.strip(), "header_row": index, "template_kind": "z-reports" if is_z else "receipts"}
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("Şablon okunamadı. Geçerli .xls veya .xlsx dosyası yükleyin.") from error


def journal_rows(documents, profile):
    if profile.get("program") not in {program["id"] for program in PROGRAMS}:
        raise ValueError("Muhasebe programını seçin.")
    accounts = STANDARD_ACCOUNTS
    chart = profile.get("chart")
    rows = []
    for sequence, document in enumerate(documents, 1):
        data = document["result"]
        if data.get("issues"):
            raise ValueError("Belgenin okuma kontrolleri tamamlanmadı; yeniden okuyun.")
        expense = document["direction"] == "expense"
        start = len(rows)
        is_z = document["kind"] == "z-reports"
        description = f"Z Raporu {data['document_no']}" if is_z else receipt_description(data)
        if data.get("extraction_version", 0) < EXTRACTION_VERSION:
            raise ValueError("Güncel okuma kontrolleri için kaynak dosya yeniden okunmalı.")
        if not data.get('seller_name') or data.get('field_sources', {}).get('seller_name'):
            raise ValueError('Firma adı kaynak belgeden doğrulanmalı; başka belgeden tamamlanan unvan aktarılamaz.')
        if not data.get("document_time") or (is_z and (not data.get("tax_office") or not (data.get("fiscal_id") or data.get("device_no")) or data.get("transaction_count") is None)):
            raise ValueError("Zorunlu belge bilgileri eksik; kaynak dosya yeniden okunmalı.")
        if document.get("chart_id", "") != profile.get("chart_id", ""):
            raise ValueError("Belge farklı firma hesap planına ait.")

        def account_for(role, rate=None, payment=None):
            if chart:
                return select_account(chart, role, rate=rate, products=[item["name"] for item in data.get("items", [])], payment=payment, evidence=data.get("bank_evidence"), expense=expense)
            code = accounts[role]
            return {"code": code, "name": ACCOUNT_NAMES[code], "bank_code": "", "card_last4": ""}

        def append(account, debit=Decimal(0), credit=Decimal(0), rate="", payment=None):
            rows.append({"date": datetime.fromisoformat(data["document_datetime"]), "voucher": str(sequence),
                         "document_no": data["document_no"], "account": account["code"], "account_name": account["name"], "description": description,
                         "document_type": "Z Raporu" if is_z else data.get("document_type", "Yazar Kasa Fişi"),
                         "document_series": data.get("document_series", ""), "seller_name": data.get("seller_name", ""),
                         "debit": debit, "credit": credit, "tax_id": data["tax_id"], "rate": rate,
                         "direction": "Gider" if expense else "Gelir", "filename": document["filename"], "currency": "TRY",
                         "bank_code": account.get("bank_code", ""), "bank_name": BANK_BY_CODE.get(account.get("bank_code"), {}).get("name", ""),
                         "payment_method": (payment or {}).get("method", ""), "card_last4": account.get("card_last4", "")})
            rows[-1].update({key: data.get(key, "") for key in ("tax_office", "document_time", "fiscal_id", "device_no", "transaction_count")})
            for key in ("total_amount", "vat_amount", "cumulative_sales", "cumulative_vat"):
                rows[-1][key] = decimal_money(data[key]) if data.get(key) not in (None, "") else None
            discount = data.get("adjustments", {}).get("discount", {}).get("amount", "")
            rows[-1]["discount_amount"] = decimal_money(discount) if discount != "" else None

        for part in data["vat_breakdown"]:
            base, tax = decimal_money(part["base"]), decimal_money(part["tax"])
            code = account_for("expense" if expense else "income", part["rate"])
            append(code, debit=base if expense else Decimal(0), credit=Decimal(0) if expense else base, rate=part["rate"])
            if tax:
                append(account_for("input_vat" if expense else "output_vat", part["rate"]), debit=tax if expense else Decimal(0), credit=Decimal(0) if expense else tax, rate=part["rate"])
        if data.get("extraction_version", 0) < 3:
            parsed, payment_issues = extract_payments(data.get("raw_text", ""), data["total_amount"], document["kind"])
            if payment_issues:
                raise ValueError("Ödeme bilgileri için belgeyi yeniden okutun.")
            payments = parsed["payment_entries"]
        else:
            payments = data.get("payment_entries", [])
        if not payments:
            raise ValueError("Ödeme dağılımı eksik. Belgeyi yeniden okutun.")
        for payment in payments:
            method, value = payment["method"], payment["amount"]
            if expense and method == "pos":
                raise ValueError("Ödemede banka kartı / kredi kartı ayrımı okunmalı.")
            if value and decimal_money(value):
                amount = decimal_money(value)
                if amount < 0:
                    raise ValueError("Negatif ödeme satırı otomatik aktarılamaz.")
                payment_key = "cash" if method == "cash" else "bank" if method == "bank_transfer" or (expense and method == "debit_card") else ("expense_" if expense else "income_") + ("card" if method in {"card", "debit_card", "pos"} else "other")
                append(account_for(payment_key, payment=payment), debit=Decimal(0) if expense else amount, credit=amount if expense else Decimal(0), payment=payment)
        if sum((row["debit"] - row["credit"] for row in rows[start:]), Decimal(0)) != 0:
            raise ValueError(f"{document['filename']}: borç ve alacak dengesi sağlanamadı.")
    return rows


def export_workbook(documents, profile, mapped=False):
    profile = resolve_profile(profile)
    kinds = {document["kind"] for document in documents}
    if len(kinds) != 1 or not kinds.issubset({"receipts", "z-reports"}):
        raise ValueError("Fiş ve Z raporu ayrı Excel dosyalarına aktarılmalı.")
    kind = kinds.pop()
    if mapped and profile["template_kind"] != kind:
        raise ValueError("Seçili şablon bu belge türüne ait değil.")
    rows = journal_rows(documents, profile)
    book = Workbook()
    sheet = book.active
    sheet.title = profile["sheet_name"] if mapped else "Fiş" if kind == "receipts" else "Z Raporu"
    columns = profile["columns"] if mapped else RECEIPT_COLUMNS if kind == "receipts" else Z_COLUMNS
    header_row = profile["header_row"] if mapped else 1

    def put(ws, row, col, value):
        cell = ws.cell(row, col, float(value) if isinstance(value, Decimal) else value)
        # Treat extracted/user-provided strings as text, never as formulas.
        if isinstance(value, str):
            cell.data_type = "s"
        return cell

    for col, spec in enumerate(columns, 1):
        put(sheet, header_row, col, spec["header"])
    for row_number, record in enumerate(rows, header_row + 1):
        for col, spec in enumerate(columns, 1):
            value = spec.get("value", "") if spec["field"] == "constant" else record.get(spec["field"], "")
            if spec["field"] == "payment_method" and value:
                from services.banking import PAYMENT_LABELS
                value = PAYMENT_LABELS.get(value, value)
            cell = put(sheet, row_number, col, value)
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=spec["field"] in {"description", "account_name"})
            if spec["field"] in {"debit", "credit", "total_amount", "vat_amount", "discount_amount", "cumulative_sales", "cumulative_vat"}:
                cell.number_format = '#,##0.00;[Red](#,##0.00);" "'
            elif spec["field"] == "date":
                cell.number_format = profile["date_format"]
            elif spec["field"] in {"account", "tax_id", "document_no", "voucher", "document_series", "bank_code", "card_last4", "fiscal_id", "device_no"}:
                cell.number_format = "@"
        wrapped = max((len(textwrap.wrap(str(record.get(spec["field"], "")), width=50 if spec["field"] == "description" else 30)) for spec in columns if spec["field"] in {"description", "account_name"}), default=1)
        sheet.row_dimensions[row_number].height = max(30, min(409, 15 * wrapped + 8))
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(max(sheet.max_row, header_row), len(columns)).coordinate}"
    for cell in sheet[header_row]:
        cell.fill = PatternFill("solid", fgColor="173D37")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 32
    for col, spec in enumerate(columns, 1):
        width = {"description": 58, "account_name": 35, "document_type": 22, "filename": 42, "document_series": 10, "blank": 2}.get(spec["field"], max(17, min(28, len(spec["header"]) + 3)))
        sheet.column_dimensions[sheet.cell(header_row, col).column_letter].width = width
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if not mapped:
        # Match the examples' summary footer; keep it outside the filtered data.
        total_row = sheet.max_row + 1
        for col, spec in enumerate(columns, 1):
            value = sum((record[spec["field"]] for record in rows), Decimal(0)) if spec["field"] in {"debit", "credit"} else len(rows) if col == 1 and kind == "receipts" else ""
            cell = put(sheet, total_row, col, value)
            cell.font = Font(name="Calibri", size=11, bold=True)
            cell.fill = PatternFill("solid", fgColor="E8F0ED")
            cell.border = Border(top=Side(style="thin", color="738B82"))
            if spec["field"] in {"debit", "credit"}:
                cell.number_format = '#,##0.00'
        sheet.row_dimensions[total_row].height = 26
        info = book.create_sheet("Aktarım Bilgisi")
        program = next(p["name"] for p in PROGRAMS if p["id"] == profile["program"])
        info_rows = [("Program", program), ("Sürüm", profile["version"]), ("Belge sayısı", len(documents)),
                     ("Hesap planı", profile.get("chart", {}).get("name", profile["account_plan"]["name"])),
                     ("Çıktı türü", "Fiş örneği sütun düzeni" if kind == "receipts" else "Z raporu örneği sütun düzeni"),
                     ("Aktarım", "Programınıza doğrudan yüklemek için kendi şablonunuzu ayarlarda eşleştirin."),
                     ("Hesap eşleşmesi", "Firmanın yüklediği hesap planındaki tekil eşleşmeler kullanıldı." if profile.get("chart") else "Standart ana hesaplar kullanıldı; firmaya özel alt hesaplar için hesap planı yükleyin.")]
        for row_index, values in enumerate(info_rows, 1):
            for col_index, value in enumerate(values, 1):
                put(info, row_index, col_index, value).alignment = Alignment(wrap_text=True, vertical="top")
            info.row_dimensions[row_index].height = 34
        info.column_dimensions["A"].width = 22
        info.column_dimensions["B"].width = 95
    from services.detail_export import add_detail_sheets
    add_detail_sheets(book, documents, profile, journal_rows)
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return output

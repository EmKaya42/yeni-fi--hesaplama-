"""Import a firm's actual chart and select only unambiguous posting accounts."""
from __future__ import annotations

import csv
import io
import re
import zipfile

from openpyxl import load_workbook

from services.banking import BANK_BY_CODE, bank_from_iban, identify_banks, normalize_iban
from services.document_extraction import folded

HEADERS = {
    "HESAPKODU": "code", "HESAPNO": "code", "HESAPNUMARASI": "code", "ACCOUNT CODE": "code",
    "HESAPADI": "name", "HESAPISMI": "name", "ACCOUNT NAME": "name",
    "IBAN": "iban", "BANKAKODU": "bank_code", "EFTKODU": "bank_code", "BANKAADI": "bank_name",
    "KARTSON4HANE": "card_last4", "SON4HANE": "card_last4", "KARTSON4": "card_last4",
    "KDVORANI": "rate", "PARABIRIMI": "currency", "DOVIZCINSI": "currency",
}
CATEGORIES = {
    "kirtasiye": (r"KIRTASIYE|OFIS SARF", r"\b(?:DEFTER|KALEM|SILGI|KAGIT|KIRTASIYE|TONER|KARTUS|ZIMBA)\b"),
    "akaryakit": (r"AKARYAKIT|YAKIT GIDER", r"\b(?:BENZIN|MOTORIN|DIZEL|AKARYAKIT|LPG)\b"),
    "temizlik": (r"TEMIZLIK", r"\b(?:DETERJAN|SABUN|CAMASIR SUYU|TEMIZLIK|YUZ[E]*Y TEMIZLEYICI)\b"),
    "yemek": (r"YEMEK|TEMSIL|AGIRLAMA", r"\b(?:YEMEK|CORBA|KEBAP|DONER|PIDE|LAHMACUN|KOFTE)\b"),
    "haberlesme": (r"HABERLESME|TELEFON|INTERNET", r"\b(?:TELEFON FATURASI|INTERNET FATURASI|HABERLESME)\b"),
}


def header_key(value):
    return re.sub(r"[^A-Z0-9]", "", folded(str(value or "")))


def compact_code(value):
    return re.sub(r"[ ._/-]", "", folded(value))


def read_chart(stream, filename):
    raw = stream.read(5 * 1024 * 1024 + 1)
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError("Hesap planı en fazla 5 MB olabilir.")
    try:
        sheets = []
        if filename.lower().endswith(".csv"):
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp1254")
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=";,\t")
            sheets = [list(csv.reader(io.StringIO(text), dialect))]
        elif raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
            import xlrd
            with xlrd.open_workbook(file_contents=raw, on_demand=True) as book:
                for sheet in book.sheets():
                    if sheet.nrows > 10020:
                        raise ValueError("Hesap planı en fazla 10.000 hesap içerebilir.")
                    sheets.append([sheet.row_values(row)[:80] for row in range(sheet.nrows)])
        else:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("Hesap planının açılmış boyutu çok büyük.")
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                for sheet in book.worksheets:
                    if (sheet.max_row or 0) > 10020:
                        raise ValueError("Hesap planı en fazla 10.000 hesap içerebilir.")
                    sheets.append(list(sheet.iter_rows(max_row=10021, max_col=min(sheet.max_column or 1, 80), values_only=True)))
            finally:
                book.close()
        aliases = {header_key(key): value for key, value in HEADERS.items()}
        candidates = []
        for rows in sheets:
            for index, row in enumerate(rows[:20]):
                mapping = {aliases[header_key(value)]: column for column, value in enumerate(row) if header_key(value) in aliases}
                if {"code", "name"}.issubset(mapping):
                    candidates.append((len(mapping), rows, index, mapping))
        if not candidates:
            raise ValueError("Hesap Kodu ve Hesap Adı sütunları bulunamadı. Firmanın hesap planını yükleyin.")
        _, rows, start, mapping = max(candidates, key=lambda candidate: candidate[0])
        accounts = {}
        for number, row in enumerate(rows[start+1:], start+2):
            record = {key: row[col] if col < len(row) else None for key, col in mapping.items()}
            code = record.get("code")
            if code in (None, ""):
                continue
            if isinstance(code, float):
                if not code.is_integer():
                    raise ValueError(f"{number}. satır: hesap kodu sayısal biçimde bozulmuş. Kaynakta metin olarak dışa aktarın.")
                code = int(code)
            code = str(code).strip()
            if re.fullmatch(r"\d{1,2}", code) or header_key(code) in {"TOPLAM", "GENELTOPLAM", "HESAPKODU"}:
                continue
            name = str(record.get("name") or "").strip()
            if not re.fullmatch(r"[1-9]\d{2}[\w ._/-]{0,61}", code) or not name or name.startswith("=") or len(name) > 200:
                raise ValueError(f"{number}. satır: hesap kodu veya hesap adı geçersiz.")
            root = code[:3]
            roles = {"770": "expense", "600": "income", "191": "input_vat", "391": "output_vat", "100": "cash", "102": "bank", "108": "income_card", "320": "expense_other", "120": "income_other"}
            role = roles.get(root, "")
            if root in {"300", "309"} and re.search(r"KREDI KART|CREDIT CARD", folded(name)):
                role = "expense_card"
            iban = normalize_iban(record["iban"]) if record.get("iban") else ""
            bank_code = str(record.get("bank_code") or "").strip()
            if bank_code:
                if re.fullmatch(r"\d+\.0", bank_code):
                    bank_code = bank_code[:-2]
                bank_code = bank_code.zfill(4)
                if bank_code not in BANK_BY_CODE:
                    raise ValueError(f"{number}. satır: banka kodu TCMB kataloğunda bulunamadı.")
            banks = identify_banks(str(record.get("bank_name") or name))
            iban_bank = bank_from_iban(iban) if iban else ""
            known = {value for value in [bank_code, iban_bank, *banks] if value}
            if len(known) > 1:
                raise ValueError(f"{number}. satır: banka adı, kodu veya IBAN birbiriyle çelişiyor.")
            bank_code = next(iter(known), "")
            last4 = str(record.get("card_last4") or "").strip()
            if last4:
                last4 = re.sub(r"\.0$", "", last4).zfill(4)
                if not re.fullmatch(r"\d{4}", last4):
                    raise ValueError(f"{number}. satır: kart son 4 hanesi geçersiz.")
            rate = record.get("rate")
            if rate not in (None, ""):
                value = str(rate).replace("%", "").replace(",", ".").strip()
                rate = float(value)
                if 0 < rate < 1:
                    rate *= 100
                if not rate.is_integer() or int(rate) not in {0, 1, 8, 10, 18, 20}:
                    raise ValueError(f"{number}. satır: KDV oranı geçersiz.")
                rate = int(rate)
            else:
                found = re.findall(r"%\s*(\d{1,2})\b|\b(\d{1,2})\s*(?:%|KDV)", folded(name))
                rates = {int(value) for group in found for value in group if value}
                rate = rates.pop() if len(rates) == 1 else None
            named_currency = re.findall(r"\b(?:USD|EUR|GBP|CHF|DOLAR|EURO)\b", folded(name))
            currency_aliases = {"DOLAR": "USD", "EURO": "EUR", "TL": "TRY", "YTL": "TRY"}
            named_currencies = {currency_aliases.get(value, value) for value in named_currency}
            inferred_currency = next(iter(named_currencies), "TRY")
            currency = folded(str(record.get("currency") or inferred_currency)).strip()
            currency = currency_aliases.get(currency, currency)
            if len(named_currencies) > 1 or (named_currencies and currency not in named_currencies):
                raise ValueError(f"{number}. satır: hesap adı ve para birimi çelişiyor.")
            categories = [key for key, (pattern, _) in CATEGORIES.items() if re.search(pattern, folded(name))] if role == "expense" else []
            value = {"code": code, "name": name, "role": role, "bank_code": bank_code, "iban": iban, "card_last4": last4, "rate": rate, "currency": currency, "categories": categories}
            key = compact_code(code)
            if key in accounts and accounts[key] != value:
                raise ValueError(f"{number}. satır: aynı hesap koduyla farklı kayıtlar var.")
            accounts[key] = value
        if not accounts or len(accounts) > 10000:
            raise ValueError("Hesap planında 1 ile 10.000 arasında geçerli hesap bulunmalı.")
        parents = set()
        for account in accounts.values():
            parts = re.split(r"[ ._/-]+", account["code"])
            parents.update(compact_code(".".join(parts[:index])) for index in range(1, len(parts)))
            if len(compact_code(account["code"])) > 3:
                parents.add(account["code"][:3])
        for key, account in accounts.items():
            account["leaf"] = key not in parents
        return list(accounts.values())
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("Hesap planı okunamadı. Geçerli XLS, XLSX veya CSV yükleyin.") from error


def select_account(chart, role, *, rate=None, products=(), payment=None, evidence=None, expense=True):
    if "by_role" not in chart:
        chart["by_role"] = {}
        for account in chart["accounts"]:
            if account["leaf"] and account["currency"] == "TRY":
                chart["by_role"].setdefault(account["role"], []).append(account)
    candidates = chart["by_role"].get(role, [])
    if rate is not None and role in {"input_vat", "output_vat", "income"}:
        exact = [account for account in candidates if account["rate"] == rate]
        candidates = exact or [account for account in candidates if account["rate"] is None]
    if role == "expense":
        categories = set()
        for product in products:
            matches = {key for key, (_, pattern) in CATEGORIES.items() if re.search(pattern, folded(product))}
            if not matches:
                categories.add("unknown")
            categories.update(matches)
        exact = [account for account in candidates if categories and categories.issubset(set(account["categories"]))]
        candidates = exact or [account for account in candidates if not account["categories"]]
    if payment is not None:
        evidence = evidence or {}
        meal_pattern = r"YEMEK KART|PLUXEE|SODEXO|MULTINET|EDENRED|SETCARD|METROPOL|TICKET"
        if payment["method"] == "meal_card":
            candidates = [account for account in chart["accounts"] if account["leaf"] and account["currency"] == "TRY"
                          and account["code"][:3] in ({"320", "329"} if expense else {"108", "120"})
                          and re.search(meal_pattern, folded(account["name"]))]
            if payment.get("provider"):
                provider = payment["provider"]
                provider_pattern = "PLUXEE|SODEXO" if provider in {"PLUXEE", "SODEXO"} else "EDENRED|TICKET" if provider.startswith("TICKET") or provider == "EDENRED" else re.escape(provider)
                candidates = [account for account in candidates if re.search(provider_pattern, folded(account["name"]).replace(" ", ""))]
        else:
            candidates = [account for account in candidates if not re.search(meal_pattern, folded(account["name"]))]
        # Printed POS bank is the merchant's acquiring bank, not proof of the
        # business credit/debit card's issuing bank.
        codes = evidence.get("issuer_codes", []) if role == "expense_card" or payment["method"] == "debit_card" else [payment.get("bank_code")] if role == "income_card" else []
        codes = [code for code in codes if code]
        last4 = evidence.get("card_last4", []) if role == "expense_card" or payment["method"] == "debit_card" else []
        if len(last4) > 1 or len(codes) > 1:
            candidates = []
        elif last4:
            candidates = [account for account in candidates if account["card_last4"] == last4[0]]
        if codes:
            candidates = [account for account in candidates if account["bank_code"] == codes[0]]
        if role == "bank" and payment["method"] == "bank_transfer":
            ibans = evidence.get("sender_ibans" if expense else "recipient_ibans", [])
            if len(ibans) > 1:
                candidates = []
            elif ibans:
                candidates = [account for account in candidates if account["iban"] == ibans[0]]
    if len(candidates) != 1:
        role_names = {"expense": "gider", "income": "gelir", "input_vat": "indirilecek KDV", "output_vat": "hesaplanan KDV", "cash": "kasa", "bank": "banka", "expense_card": "kredi kartı", "income_card": "POS tahsilat", "expense_other": "satıcı", "income_other": "alıcı"}
        reason = "birden fazla uygun hesap var" if candidates else "uygun hesap bulunamadı"
        raise ValueError(f"{role_names.get(role, role).capitalize()} hesabı eşleşmedi: {reason}. Firma hesap planındaki banka, kart son 4 hanesi, KDV ve hesap adlarını kontrol edin.")
    return candidates[0]

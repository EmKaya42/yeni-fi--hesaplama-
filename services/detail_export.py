"""Turkish detail sheets appended after the existing program import sheet."""
from datetime import datetime, time
from decimal import Decimal
import textwrap

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.banking import BANK_BY_CODE, PAYMENT_LABELS
from services.document_details import ADJUSTMENT_LABELS
from services.document_extraction import decimal_money


def amount(value):
    return decimal_money(value) if value not in (None, "") else None


def add_detail_sheets(book, documents, profile, journal_reader):
    from services.accounting import PROGRAMS
    program_name = next(p["name"] for p in PROGRAMS if p["id"] == profile["program"])
    def table(title, headers, rows, money_columns=(), number_columns=()):
        sheet = book.create_sheet(title)
        widths = [42 if header in {"Ürün / hizmet adı", "İşletme adı / unvanı", "Kaynak dosya", "Hesap adı"} else max(18, min(30, len(header) + 2)) for header in headers]
        for row_no, values in enumerate([headers, *rows], 1):
            for col_no, value in enumerate(values, 1):
                cell = sheet.cell(row_no, col_no, float(value) if isinstance(value, Decimal) else value)
                if isinstance(value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                if row_no == 1:
                    cell.fill = PatternFill("solid", fgColor="173D37")
                    cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
                else:
                    cell.font = Font(name="Calibri", size=11)
                    if col_no in money_columns:
                        cell.number_format = '#,##0.00;[Red](#,##0.00);"0,00"'
                    elif col_no in number_columns:
                        cell.number_format = "0.######"
                    elif isinstance(value, datetime):
                        cell.number_format = "dd.mm.yyyy"
                    elif isinstance(value, time):
                        cell.number_format = "hh:mm:ss"
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            wrapped = max((len(textwrap.wrap(str(value or ""), width=max(10, int(widths[index]) - 2))) for index, value in enumerate(values)), default=1)
            sheet.row_dimensions[row_no].height = min(409, max(42 if row_no == 1 else 34, 15 * wrapped + 8))
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        return sheet

    is_z = documents[0]["kind"] == "z-reports"
    info, items, vat, payments, adjustments = [], [], [], [], []
    for document in documents:
        data = document["result"]
        date = datetime.fromisoformat(data["document_datetime"][:10])
        clock = time.fromisoformat(data["document_time"]) if data.get("document_time") else None
        base = [document["filename"], data["document_no"], data.get("seller_name", "")]
        adjustment = data.get("adjustments", {})
        info.append(base + [data.get("tax_office", ""), data["tax_id"], date, clock,
                           data.get("fiscal_id", ""), data.get("device_no", ""),
                           amount(data["total_amount"]), amount(data["vat_amount"]),
                           amount(adjustment.get("discount", {}).get("amount")),
                           amount(adjustment.get("cancellation", {}).get("amount")), adjustment.get("cancellation", {}).get("count"),
                           amount(adjustment.get("refund", {}).get("amount")), adjustment.get("refund", {}).get("count"),
                           data.get("transaction_count"), amount(data.get("cumulative_sales")), amount(data.get("cumulative_vat")),
                           profile.get("chart", {}).get("name", "Standart hesaplar"),
                           program_name, data.get('field_sources', {}).get('seller_name', {}).get('filename', ''),
                           '\n'.join(data.get('notes', []))])
        for index, item in enumerate(data.get("items", []), 1):
            items.append(base + [index, item["name"], Decimal(item["quantity"]) if item.get("quantity") else None,
                                item.get("unit", ""), amount(item.get("unit_price")), amount(item["amount"]), item.get("rate")])
        for part in data["vat_breakdown"]:
            vat.append(base + [part["rate"], amount(part["base"]) + amount(part["tax"]), amount(part["base"]), amount(part["tax"])])
        journal_payments = iter(row for row in journal_reader([document], profile) if row["payment_method"])
        for payment in data.get("payment_entries", []):
            matched = next(journal_payments, {}) if decimal_money(payment["amount"]) else {}
            bank = BANK_BY_CODE.get(payment.get("bank_code"), {})
            payments.append(base + [PAYMENT_LABELS[payment["method"]], payment.get("provider", ""), amount(payment["amount"]),
                                    bank.get("name", ""), bank.get("code", ""), matched.get("account", ""), matched.get("account_name", ""),
                                    matched.get("bank_name", ""), matched.get("bank_code", ""), matched.get("card_last4", "")])
        for key, value in adjustment.items():
            adjustments.append(base + [ADJUSTMENT_LABELS[key], amount(value.get("amount")), value.get("count")])
    identifiers = ["Kaynak dosya", "Z raporu numarası" if is_z else "Fiş / belge numarası", "İşletme adı / unvanı"]
    table("Belge Bilgileri", identifiers + ["Vergi dairesi", "VKN / TCKN", "Tarih", "Saat", "Mali sicil numarası", "Cihaz numarası",
          "KDV dâhil toplam satış" if is_z else "Genel toplam", "Toplam KDV", "İndirim tutarı", "İptal tutarı", "İptal adedi", "İade tutarı", "İade adedi",
          "Fiş / işlem adedi", "Kümülatif satış", "Kümülatif KDV", "Firma hesap planı", "Muhasebe programı", "Otomatik firma adı kaynağı", "Okuma notları"], info,
          money_columns=(10, 11, 12, 13, 15, 18, 19), number_columns=(14, 16, 17))
    if not is_z:
        table("Fiş Kalemleri", identifiers + ["Kalem sırası", "Ürün / hizmet adı", "Miktar", "Birim", "Birim fiyat", "Kalem tutarı", "KDV oranı (%)"], items,
              money_columns=(8, 9), number_columns=(4, 6, 10))
    table("KDV Dağılımı", identifiers + ["KDV oranı (%)", "KDV dâhil satış tutarı", "KDV hariç matrah", "KDV tutarı"], vat, money_columns=(5, 6, 7), number_columns=(4,))
    table("Ödeme Dağılımı", identifiers + ["Ödeme şekli", "Yemek kartı kuruluşu", "Tutar", "Belgede basılı banka / POS bankası", "Basılı banka EFT kodu",
          "Muhasebe hesap kodu", "Hesap adı", "Eşleşen hesap bankası", "Eşleşen banka EFT kodu", "Kart son 4 hane"], payments, money_columns=(6,))
    table("İndirim İptal İade", identifiers + ["İşlem türü", "Tutar", "Adet"], adjustments, money_columns=(5,), number_columns=(6,))
